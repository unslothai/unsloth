# SPDX-License-Identifier: AGPL-3.0-only
# Copyright 2026-present the Unsloth AI Inc. team. All rights reserved. See /studio/LICENSE.AGPL-3.0

"""Office, OpenDocument, e-book, email and RTF files: parsing, upload and indexing."""

import datetime as dt
import io
import re
import struct
import zipfile
from email.message import EmailMessage
from xml.sax.saxutils import escape
from pathlib import Path

import pytest
from fastapi import UploadFile

from core.rag import cfb, config, ingestion, office_formats, parsers, store
from routes.rag import _save_upload
from storage import rag_db

from ._compound_file import compound_file

S = 'xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
P = 'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
REL = 'xmlns="http://schemas.openxmlformats.org/package/2006/relationships"'
ODF = (
    'xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
    'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" '
    'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
    'xmlns:draw="urn:oasis:names:tc:opendocument:xmlns:drawing:1.0" '
    'xmlns:presentation="urn:oasis:names:tc:opendocument:xmlns:presentation:1.0"'
)


def _zip(path, members):
    with zipfile.ZipFile(path, "w") as z:
        for name, data in members.items():
            z.writestr(name, data)
    return path


def _text(path):
    return "\n".join(page.text for page in parsers.parse(str(path)))


def build_xlsx(path):
    serial = (dt.date(2026, 3, 31) - dt.date(1899, 12, 30)).days
    return _zip(
        path,
        {
            "xl/workbook.xml": f'<workbook {S}><sheets><sheet name="Revenue" sheetId="1" r:id="rId1"/></sheets></workbook>',
            "xl/_rels/workbook.xml.rels": f'<Relationships {REL}><Relationship Id="rId1" Target="worksheets/sheet1.xml"/></Relationships>',
            "xl/sharedStrings.xml": f"<sst {S}><si><t>Region</t></si>"
            "<si><r><t>Zebra</t></r><r><t>marker</t></r><rPh><t>ignored</t></rPh></si></sst>",
            "xl/styles.xml": f'<styleSheet {S}><cellXfs><xf numFmtId="0"/><xf numFmtId="14"/></cellXfs></styleSheet>',
            "xl/worksheets/sheet1.xml": f"<worksheet {S}><sheetData>"
            '<row r="1"><c r="A1" t="s"><v>0</v></c><c r="B1" t="inlineStr"><is><t>Q1</t></is></c></row>'
            f'<row r="2"><c r="A2" t="s"><v>1</v></c><c r="B2"><v>1200.5</v></c><c r="C2" s="1"><v>{serial}</v></c>'
            '<c r="D2" t="b"><v>1</v></c><c r="F2" t="str"><v>formula text</v></c></row>'
            "</sheetData></worksheet>",
        },
    )


def build_pptx(path):
    run = lambda text: f"<a:p><a:r><a:t>{text}</a:t></a:r></a:p>"
    return _zip(
        path,
        {
            "ppt/presentation.xml": f'<p:presentation {P}><p:sldIdLst><p:sldId id="256" r:id="rId2"/></p:sldIdLst></p:presentation>',
            "ppt/_rels/presentation.xml.rels": f'<Relationships {REL}><Relationship Id="rId2" Target="slides/slide1.xml"/></Relationships>',
            "ppt/slides/slide1.xml": f"<p:sld {P}><p:cSld><p:spTree>"
            f"<p:sp><p:txBody>{run('Quarterly review')}</p:txBody></p:sp>"
            "<p:graphicFrame><a:graphic><a:graphicData><a:tbl>"
            f"<a:tr><a:tc><a:txBody>{run('Region')}</a:txBody></a:tc><a:tc><a:txBody>{run('Q1')}</a:txBody></a:tc></a:tr>"
            f"<a:tr><a:tc><a:txBody>{run('North')}</a:txBody></a:tc><a:tc><a:txBody>{run('1200')}</a:txBody></a:tc></a:tr>"
            "</a:tbl></a:graphicData></a:graphic></p:graphicFrame>"
            f"<p:sp><p:txBody>{run('Revenue increased')}</p:txBody></p:sp>"
            "</p:spTree></p:cSld></p:sld>",
            "ppt/slides/_rels/slide1.xml.rels": f'<Relationships {REL}><Relationship Id="rId1" Target="../notesSlides/notesSlide1.xml"/></Relationships>',
            "ppt/notesSlides/notesSlide1.xml": f"<p:notes {P}><p:cSld><p:spTree>"
            '<p:sp><p:nvSpPr><p:cNvPr id="5" name="Slide Number"/><p:cNvSpPr/><p:nvPr><p:ph type="sldNum" idx="5"/></p:nvPr></p:nvSpPr>'
            '<p:txBody><a:p><a:fld id="{1}" type="slidenum"><a:t>1</a:t></a:fld></a:p></p:txBody></p:sp>'
            f"<p:sp><p:txBody>{run('2026')}{run('Speaker notes zebramarker')}</p:txBody></p:sp>"
            "</p:spTree></p:cSld></p:notes>",
        },
    )


def _odf(path, body):
    suffix = Path(path).suffix
    kind = {".odt": "text", ".ods": "spreadsheet", ".odp": "presentation"}.get(suffix)
    if kind is None:
        kind = {".ott": "text", ".ots": "spreadsheet", ".otp": "presentation"}[suffix] + "-template"
    return _zip(
        path,
        {
            "mimetype": f"application/vnd.oasis.opendocument.{kind}",
            "content.xml": f"<office:document-content {ODF}><office:body>{body}</office:body></office:document-content>",
        },
    )


def build_odt(path):
    return _odf(
        path,
        "<office:text><text:h>Quarterly report</text:h>"
        '<text:p>Revenue<text:s text:c="2"/>increased<text:note><text:note-body><text:p>see appendix</text:p></text:note-body></text:note>.</text:p>'
        "<table:table><table:table-row><table:table-cell><text:p>Zebramarker</text:p></table:table-cell>"
        "<table:table-cell><text:p>1200</text:p></table:table-cell></table:table-row></table:table></office:text>",
    )


def build_ods(path):
    return _odf(
        path,
        '<office:spreadsheet><table:table table:name="Revenue">'
        "<table:table-row><table:table-cell><text:p>Region</text:p></table:table-cell>"
        '<table:table-cell table:number-columns-repeated="2"><text:p>Q</text:p></table:table-cell>'
        '<table:table-cell table:number-columns-repeated="16000"/></table:table-row>'
        '<table:table-row table:number-rows-repeated="1048000"><table:table-cell table:number-columns-repeated="16000"/></table:table-row>'
        "<table:table-row><table:table-cell><text:p>Zebramarker</text:p></table:table-cell></table:table-row>"
        "</table:table></office:spreadsheet>",
    )


def build_odp(path):
    return _odf(
        path,
        "<office:presentation>"
        "<draw:page><draw:frame><draw:text-box><text:p>First slide</text:p></draw:text-box></draw:frame>"
        "<presentation:notes><draw:frame><draw:text-box><text:p>zebramarker notes</text:p></draw:text-box></draw:frame></presentation:notes></draw:page>"
        "<draw:page><draw:frame><draw:text-box><text:p>Second slide revenue</text:p></draw:text-box></draw:frame></draw:page>"
        "</office:presentation>",
    )


def build_epub(path):
    page = (
        lambda title,
        text: f"<html><head><title>{title}</title></head><body><p>{text}</p></body></html>"
    )
    return _zip(
        path,
        {
            "mimetype": "application/epub+zip",
            "META-INF/container.xml": '<container xmlns="urn:oasis:names:tc:opendocument:xmlns:container">'
            '<rootfiles><rootfile full-path="OEBPS/content.opf"/></rootfiles></container>',
            "OEBPS/content.opf": '<package xmlns="http://www.idpf.org/2007/opf"><manifest>'
            '<item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/>'
            '<item id="c1" href="text/Chapter%201.xhtml#start" media-type="application/xhtml+xml"/>'
            '<item id="c2" href="text/c2.xhtml" media-type="application/xhtml+xml"/>'
            '</manifest><spine><itemref idref="nav"/><itemref idref="c2"/><itemref idref="c1"/></spine></package>',
            "OEBPS/nav.xhtml": page("Contents", "navmarker"),
            "OEBPS/text/Chapter 1.xhtml": page("One", "Zebramarker ends the book."),
            "OEBPS/text/c2.xhtml": page("Two", "Revenue increased."),
        },
    )


def build_eml(path):
    message = EmailMessage()
    message["From"] = "Ann <ann@example.com>"
    message["To"] = "bob@example.com"
    message["Subject"] = "Q1 numbers résumé"
    message.set_content("Revenue increased.\nZebramarker in the body.", cte = "base64")
    message.add_alternative("<p>html copy</p>", subtype = "html")
    message.add_attachment(b"x,y\n", maintype = "text", subtype = "csv", filename = "data.csv")
    path.write_bytes(bytes(message))
    return path


def build_mhtml(path):
    # As Chrome saves a page: multipart/related, the page plus its resources.
    message = EmailMessage()
    message["From"] = "<Saved by Blink>"
    message["Subject"] = "Quarterly report"
    message.set_content(
        "<html><body><p>Zebramarker on the page.</p><img src='logo.png'></body></html>",
        subtype = "html",
    )
    message.make_related()
    message.add_related(b"\x89PNG\r\n\x1a\n", maintype = "image", subtype = "png", cid = "<logo>")
    path.write_bytes(bytes(message))
    return path


def build_xhtml(path):
    path.write_text(
        '<?xml version="1.0" encoding="utf-8"?><html xmlns="http://www.w3.org/1999/xhtml">'
        "<head><title>Report</title><style>p{}</style></head><body><p>Zebramarker résumé</p></body></html>",
        encoding = "utf-8",
    )
    return path


_WORD_TYPES = {
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml",
    ".docm": "application/vnd.ms-word.document.macroEnabled.main+xml",
    ".dotx": "application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml",
    ".dotm": "application/vnd.ms-word.template.macroEnabledTemplate.main+xml",
}


def build_word(path):
    import docx

    document = docx.Document()
    document.add_paragraph("Quarterly report")
    document.add_table(rows = 1, cols = 2).rows[0].cells[0].text = "Zebramarker"
    buffer = io.BytesIO()
    document.save(buffer)
    with zipfile.ZipFile(buffer) as source, zipfile.ZipFile(path, "w") as target:
        for info in source.infolist():
            data = source.read(info)
            if info.filename == "[Content_Types].xml":
                data = data.replace(
                    _WORD_TYPES[".docx"].encode(), _WORD_TYPES[path.suffix].encode()
                )
            target.writestr(info, data)
    return path


def build_rtf(path):
    path.write_bytes(
        rb"{\rtf1\ansi\ansicpg1252\uc1{\fonttbl{\f0 Arial;}}{\colortbl;\red0\green0\blue0;}"
        rb"{\*\generator Writer;}\f0 Quarterly report\par Caf\'e9 \u8364? revenue increased\par"
        rb"{\field{\*\fldinst HYPERLINK hiddenmarker}{\fldrslt link text}}\par"
        rb"\trowd Zebramarker\cell 1200\cell\row}"
    )
    return path


def build_doc(path, *, encrypted = False):
    compressed = "Quarterly report\r".encode("cp1252")
    wide = ("Café 東京 \x13 HYPERLINK x \x14link\x15 Zebramarker\x071200\x07\x07").encode(
        "utf-16-le"
    )
    word = bytearray(0x1800)
    flags = 0x0200 | (0x0100 if encrypted else 0)
    struct.pack_into("<HH", word, 0, 0xA5EC, 0xC1)
    struct.pack_into("<H", word, 0x0A, flags)
    struct.pack_into("<H", word, 32, 14)  # csw
    struct.pack_into("<H", word, 62, 22)  # cslw
    ccp = len(compressed) + len(wide) // 2
    footnote, endnote, text_box = "\x02 Footnote detail\r", "\x02 Endnote source\r", "Boxed text\r"
    struct.pack_into("<i", word, 64 + 3 * 4, ccp)  # ccpText
    struct.pack_into("<i", word, 64 + 4 * 4, len(footnote))  # ccpFtn
    struct.pack_into("<i", word, 64 + 8 * 4, len(endnote))  # ccpEdn
    struct.pack_into("<i", word, 64 + 9 * 4, len(text_box))  # ccpTxbx
    struct.pack_into("<H", word, 152, 93)  # cbRgFcLcb
    stories = (footnote + endnote + text_box).encode("utf-16-le")
    word[0x1000 : 0x1000 + len(compressed)] = compressed
    word[0x1400 : 0x1400 + len(wide)] = wide
    word[0x1600 : 0x1600 + len(stories)] = stories
    cps = [0, len(compressed), ccp, ccp + len(stories) // 2]
    pcds = [(0x2000 | 0x40000000), 0x1400, 0x1600]
    plc = struct.pack("<4I", *cps) + b"".join(struct.pack("<HIH", 0, fc, 0) for fc in pcds)
    clx = b"\x01" + struct.pack("<H", 2) + b"\0\0" + b"\x02" + struct.pack("<I", len(plc)) + plc
    struct.pack_into("<II", word, 154 + 33 * 8, 0, len(clx))  # fcClx, lcbClx
    path.write_bytes(compound_file({("WordDocument",): bytes(word), ("1Table",): clx}))
    return path


def _record(kind, body):
    return struct.pack("<HH", kind, len(body)) + body


def build_xls(path):
    def short(text, size = 1):
        return struct.pack("<B" if size == 1 else "<H", len(text)) + b"\0" + text.encode("latin-1")

    # SST: "Region", then a string whose characters continue, wide, in a CONTINUE record.
    sst = struct.pack("<II", 2, 2) + struct.pack("<HB", 6, 0) + b"Region"
    sst += struct.pack("<HB", 12, 0) + b"Zebra"
    cont = b"\x01" + "marker€".encode("utf-16-le")
    xf = lambda fmt: _record(0x00E0, struct.pack("<HH", 0, fmt) + b"\0" * 16)
    globals_ = (
        _record(0x0809, b"\0" * 16)
        + _record(0x0085, b"\0" * 4 + b"\0\0" + short("Revenue"))
        + _record(0x00FC, sst)
        + _record(0x003C, cont)
        + _record(0x041E, struct.pack("<H", 164) + short("yyyy-mm-dd", 2))
        + xf(0)
        + xf(14)
        + xf(164)
        + _record(0x000A, b"")
    )
    serial = (dt.date(2026, 3, 31) - dt.date(1899, 12, 30)).days
    rk_int = lambda n: ((n << 2) | 2) & 0xFFFFFFFF
    sheet = (
        _record(0x0809, b"\0" * 16)
        + _record(0x00FD, struct.pack("<HHHI", 0, 0, 0, 0))
        + _record(0x0203, struct.pack("<HHHd", 0, 1, 0, 1200.5))
        + _record(0x00FD, struct.pack("<HHHI", 1, 0, 0, 1))
        + _record(0x027E, struct.pack("<HHHI", 1, 1, 0, rk_int(-7)))
        + _record(
            0x00BD,
            struct.pack("<HH", 2, 0)
            + struct.pack("<HI", 0, rk_int(3))
            + struct.pack("<HI", 0, rk_int(4))
            + struct.pack("<H", 1),
        )
        + _record(0x0809, b"\0" * 16)  # embedded chart
        + _record(0x000A, b"")
        + _record(0x0006, struct.pack("<HHH", 3, 0, 0) + b"\0" * 6 + b"\xff\xff" + b"\0" * 6)
        + _record(0x0207, short("formula text", 2))
        + _record(0x0205, struct.pack("<HHHBB", 3, 1, 0, 1, 0))
        + _record(0x0203, struct.pack("<HHHd", 4, 0, 2, serial))
        + _record(0x027E, struct.pack("<HHHI", 4, 1, 1, rk_int(serial + 1)))
        + _record(0x000A, b"")
    )
    stream = bytearray(globals_ + sheet)
    struct.pack_into("<I", stream, 4 + 16 + 4, len(globals_))  # BOUNDSHEET lbPlyPos
    path.write_bytes(compound_file({("Workbook",): bytes(stream)}))
    return path


def _ppt_atom(
    kind,
    body,
    inst = 0,
):
    return struct.pack("<HHI", inst << 4, kind, len(body)) + body


def _ppt_container(
    kind,
    children,
    inst = 0,
):
    body = b"".join(children)
    return struct.pack("<HHI", (inst << 4) | 0xF, kind, len(body)) + body


def build_ppt(path, *, from_slides = False):
    chars = lambda text: _ppt_atom(0x0FA0, text.encode("utf-16-le"))
    persist = _ppt_atom(0x03F3, b"\0" * 20)
    if from_slides:
        slides = [
            _ppt_container(
                0x03EE, [_ppt_container(0x040C, [_ppt_container(0xF00D, [chars(text)])])]
            )
            for text in ("Drawn slide", "Zebramarker drawn")
        ]
        document = _ppt_container(0x03E8, slides)
    else:
        document = _ppt_container(
            0x03E8,
            [
                _ppt_container(
                    0x0FF0,
                    [
                        persist,
                        chars("Quarterly review\rRevenue up"),
                        persist,
                        _ppt_atom(0x0FA8, b"Zebramarker slide"),
                    ],
                    inst = 0,
                ),
                _ppt_container(
                    0x0FF0, [persist, chars("Click to edit Master title style")], inst = 1
                ),
            ],
        )
    path.write_bytes(compound_file({("PowerPoint Document",): document}))
    return path


# PidTagRtfCompressed of "{\\rtf1\\ansi\\ansicpg1252\\pard Zebramarker in the RTF body\\par repeat repeat repeat}".
RTF_COMPRESSED = bytes.fromhex(
    "49000000520000004c5a4675a3e8a93d03000a007263706731323502320af3205a656272612900c0726b04902"
    "00b80207444686507f054462006e064c6790aa30970706561054010cb027d1210"
)


def build_ansi_msg(path):
    """String8 properties in code page 1251 and only a compressed RTF body."""
    props = b"\0" * 32 + struct.pack("<IIQ", 0x3FFD0003, 0, 1251)
    streams = {
        ("__substg1.0_0037001E",): "Привет".encode("cp1251"),
        ("__substg1.0_0C1A001E",): "Иван".encode("cp1251"),
        ("__substg1.0_10090102",): RTF_COMPRESSED,
        ("__properties_version1.0",): props,
    }
    path.write_bytes(compound_file(streams))
    return path


def build_saved_ppt(path, *, broken_link = False):
    """Saved twice: the second save replaces the slide, adds notes and leaves the old slide behind."""
    chars = lambda text: _ppt_atom(0x0FA0, text.encode("utf-16-le"))
    textbox = lambda *atoms: _ppt_container(0xF00D, list(atoms))
    drawing = lambda *boxes: _ppt_container(0x040C, [_ppt_container(0xF002, list(boxes))])
    persist = lambda ref, sheet_id: _ppt_atom(0x03F3, struct.pack("<5I", ref, 0, 1, sheet_id, 0))
    stale = _ppt_container(0x03EE, [drawing(textbox(chars("Stale slide text")))])
    slide = _ppt_container(
        0x03EE,
        [
            drawing(
                textbox(_ppt_atom(0x0F9E, struct.pack("<I", 0))),
                textbox(chars("Text box zebramarker")),
            )
        ],
    )
    notes = _ppt_container(
        0x03F0,
        [
            _ppt_atom(0x03F1, struct.pack("<IHH", 256, 0, 0)),
            drawing(textbox(chars("Speaker notes")), textbox(chars("*"))),
        ],
    )
    document = _ppt_container(
        0x03E8,
        [
            _ppt_container(
                0x0FF0,
                [persist(2, 256), _ppt_atom(0x0F9F, b"\0" * 4), chars("Title from list")],
                inst = 0,
            ),
            _ppt_container(0x0FF0, [persist(3, 257)], inst = 2),
        ],
    )
    stream = bytearray()

    def put(record):
        stream.extend(record)
        return len(stream) - len(record)

    def directory(entries):
        body = b"".join(struct.pack("<II", (1 << 20) | pid, offset) for pid, offset in entries)
        return put(_ppt_atom(0x1772, body))

    def user_edit(last, directory_offset):
        offset = len(stream)
        last = offset if broken_link and last else last
        return put(
            _ppt_atom(0x0FF5, struct.pack("<IHBBIIII", 0, 0, 0, 3, last, directory_offset, 1, 4))
        )

    stale_at, doc_at = put(stale), put(document)
    first = user_edit(0, directory([(1, doc_at), (2, stale_at)]))
    slide_at, notes_at = put(slide), put(notes)
    current = user_edit(first, directory([(2, slide_at), (3, notes_at)]))
    current_user = _ppt_atom(0x0FF6, struct.pack("<III", 20, 0xE391C05F, current))
    path.write_bytes(
        compound_file({("PowerPoint Document",): bytes(stream), ("Current User",): current_user})
    )
    return path


def build_msg(path):
    utf16 = lambda text: text.encode("utf-16-le")
    filetime = int(
        (dt.datetime(2026, 10, 6, 10, 0) - dt.datetime(1601, 1, 1)).total_seconds() * 10**7
    )
    props = b"\0" * 32 + struct.pack("<IIQ", (0x0039 << 16) | 0x0040, 0, filetime)
    streams = {
        ("__substg1.0_0037001F",): utf16("Q1 numbers"),
        ("__substg1.0_0C1A001F",): utf16("Ann"),
        ("__substg1.0_0C1F001F",): utf16("ann@example.com"),
        ("__substg1.0_0E04001F",): utf16("Bob"),
        ("__substg1.0_1000001F",): utf16("Revenue increased.\r\nZebramarker in the body."),
        ("__properties_version1.0",): props,
        ("__attach_version1.0_#00000000", "__substg1.0_3707001F"): utf16("data.csv"),
    }
    path.write_bytes(compound_file(streams))
    return path


BUILDERS = {
    ".doc": build_doc,
    ".xls": build_xls,
    ".ppt": build_ppt,
    ".msg": build_msg,
    ".eml": build_eml,
    ".rtf": build_rtf,
    ".epub": build_epub,
    **dict.fromkeys((".docm", ".dotx", ".dotm"), build_word),
    **dict.fromkeys((".xlsx", ".xlsm", ".xltx", ".xltm"), build_xlsx),
    **dict.fromkeys((".pptx", ".pptm", ".potx", ".potm", ".ppsx", ".ppsm"), build_pptx),
    **dict.fromkeys((".odt", ".ott"), build_odt),
    **dict.fromkeys((".ods", ".ots"), build_ods),
    **dict.fromkeys((".odp", ".otp"), build_odp),
    **dict.fromkeys((".mht", ".mhtml"), build_mhtml),
    **dict.fromkeys((".xhtml", ".xht"), build_xhtml),
}


def test_every_document_type_has_a_fixture():
    assert set(BUILDERS) == config.DOCUMENT_UPLOAD_EXTS
    assert config.DOCUMENT_UPLOAD_EXTS <= config.UPLOAD_EXTS


def test_document_exts_mirror_the_frontend_accept_list():
    source = (
        Path(__file__).resolve().parents[2] / "frontend/src/features/rag/types/rag.ts"
    ).read_text(encoding = "utf-8")
    accept = re.search(r'RAG_DOCUMENT_UPLOAD_ACCEPT =\s*"([^"]+)"', source).group(1)
    assert set(accept.split(",")) == config.DOCUMENT_UPLOAD_EXTS


def test_xlsx_reads_shared_inline_dates_and_gaps(tmp_path):
    text = _text(build_xlsx(tmp_path / "book.xlsx"))
    assert (
        text
        == "Sheet: Revenue\nRegion | Q1\nZebramarker | 1200.5 | 2026-03-31 | TRUE |  | formula text"
    )


def test_pptx_reads_slides_in_order_with_tables_and_notes(tmp_path):
    pages = parsers.parse(str(build_pptx(tmp_path / "deck.pptx")))
    assert [(p.page_number, p.text) for p in pages] == [
        (
            1,
            "Quarterly review\nRegion | Q1\nNorth | 1200\nRevenue increased\nNotes:\n2026\nSpeaker notes zebramarker",
        )
    ]


def test_odt_reads_paragraphs_spaces_notes_and_tables(tmp_path):
    assert _text(build_odt(tmp_path / "memo.odt")) == (
        "Quarterly report\nRevenue  increased [see appendix].\nZebramarker | 1200"
    )


def test_ods_does_not_expand_empty_repeats(tmp_path):
    assert _text(build_ods(tmp_path / "sheet.ods")) == "Sheet: Revenue\nRegion | Q | Q\nZebramarker"


def test_ods_keeps_the_width_of_empty_runs_before_content(tmp_path):
    path = _odf(
        tmp_path / "gaps.ods",
        '<office:spreadsheet><table:table table:name="S"><table:table-row>'
        "<table:table-cell><text:p>A</text:p></table:table-cell>"
        '<table:table-cell table:number-columns-repeated="4"/>'
        "<table:table-cell><text:p>F</text:p></table:table-cell>"
        '<table:table-cell table:number-columns-repeated="16000"/>'
        "</table:table-row></table:table></office:spreadsheet>",
    )
    assert _text(path) == "Sheet: S\nA |  |  |  |  | F"


def test_ods_reads_typed_values_without_text(tmp_path):
    path = _odf(
        tmp_path / "values.ods",
        '<office:spreadsheet><table:table table:name="S"><table:table-row>'
        '<table:table-cell office:value-type="float" office:value="1200"/>'
        '<table:table-cell office:value-type="date" office:date-value="2026-03-31"/>'
        '<table:table-cell office:value-type="boolean" office:boolean-value="true"/>'
        '<table:table-cell office:value-type="string" office:string-value="zebramarker"/>'
        '<table:table-cell office:value-type="float" office:value="7"><text:p>7.00</text:p></table:table-cell>'
        "</table:table-row></table:table></office:spreadsheet>",
    )
    assert _text(path) == "Sheet: S\n1200 | 2026-03-31 | TRUE | zebramarker | 7.00"


# Dates, times and durations match what openpyxl reads back from the same cells.
@pytest.mark.parametrize(
    "serial, number_format, expected",
    [
        (0.5, "h:mm AM/PM", "12:00:00"),
        (1.5, "[h]:mm:ss", "36:00:00"),
        (1.5, "h:mm", "1900-01-01 12:00:00"),
        (1, "yyyy-mm-dd", "1900-01-01"),
        (59, "yyyy-mm-dd", "1900-02-28"),
        (61, "yyyy-mm-dd", "1900-03-01"),
        (0.25, "mm-dd-yy", "06:00:00"),
        (3, '"day" 0', "day 3"),
        (0.25, "0%", "25%"),
        (0.125, "0.0%", "12.5%"),
        (123, "000000", "000123"),
        (-1234.5, "#,##0.00", "-1,234.50"),
        (-1234.5, "#,##0.00;(#,##0.00)", "(1,234.50)"),
        (12345.678, "0.00E+00", "1.23E+04"),
        (12345.678, "##0.0E+0", "12.3E+3"),
        (1234.5, '"$"#,##0', "$1,235"),
        (1234.5, "[$\u20ac-407]#,##0.00", "\u20ac1,234.50"),
        (0.5, "#.##", ".5"),
        (1.234, "0.0#", "1.23"),
        (1500000, '#,##0,,"M"', "2M"),
        (1234.5, "General", "1234.5"),
        (5, "[>100]0;0.0", "5"),
        pytest.param(1234.5, "0" * 5000, "1234.5", id = "oversized-format-left-as-stored"),
    ],
)
def test_xlsx_shows_numbers_as_formatted(tmp_path, serial, number_format, expected):
    path = _zip(
        tmp_path / "times.xlsx",
        {
            "xl/workbook.xml": f'<workbook {S}><sheets><sheet name="T" sheetId="1" r:id="rId1"/></sheets></workbook>',
            "xl/_rels/workbook.xml.rels": f'<Relationships {REL}><Relationship Id="rId1" Target="worksheets/sheet1.xml"/></Relationships>',
            "xl/styles.xml": f'<styleSheet {S}><numFmts><numFmt numFmtId="164" formatCode="{escape(number_format, {chr(34): "&quot;"})}"/></numFmts>'
            '<cellXfs><xf numFmtId="0"/><xf numFmtId="164"/></cellXfs></styleSheet>',
            "xl/worksheets/sheet1.xml": f'<worksheet {S}><sheetData><row r="1"><c r="A1" s="1"><v>{serial}</v></c></row></sheetData></worksheet>',
        },
    )
    assert _text(path) == f"Sheet: T\n{expected}"


def test_xlsx_finds_the_workbook_through_package_relationships(tmp_path):
    path = _zip(
        tmp_path / "moved.xlsx",
        {
            "_rels/.rels": f'<Relationships {REL}><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="book/main.xml"/></Relationships>',
            "book/main.xml": f'<workbook {S}><sheets><sheet name="M" sheetId="1" r:id="rId1"/></sheets></workbook>',
            "book/_rels/main.xml.rels": f'<Relationships {REL}><Relationship Id="rId1" Target="sheet.xml"/>'
            '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings" Target="strings.xml"/></Relationships>',
            "book/strings.xml": f"<sst {S}><si><t>zebramarker</t></si></sst>",
            "book/sheet.xml": f'<worksheet {S}><sheetData><row r="1"><c r="A1" t="s"><v>0</v></c></row></sheetData></worksheet>',
        },
    )
    assert _text(path) == "Sheet: M\nzebramarker"


def test_xml_parts_are_bounded_by_element_count(tmp_path):
    row = "<c/>" * 200_001  # one streamed row
    sheet_bomb = _zip(
        tmp_path / "row.xlsx",
        {
            "xl/workbook.xml": f'<workbook {S}><sheets><sheet name="B" sheetId="1" r:id="rId1"/></sheets></workbook>',
            "xl/_rels/workbook.xml.rels": f'<Relationships {REL}><Relationship Id="rId1" Target="worksheets/sheet1.xml"/></Relationships>',
            "xl/worksheets/sheet1.xml": f"<worksheet {S}><sheetData><row>{row}</row></sheetData></worksheet>",
        },
    )
    tree_bomb = _zip(
        tmp_path / "tree.xlsx",
        {"xl/workbook.xml": f"<workbook {S}>" + "<x/>" * 2_000_001 + "</workbook>"},
    )
    deep = _zip(
        tmp_path / "deep.ods",
        {
            "mimetype": "application/vnd.oasis.opendocument.spreadsheet",
            "content.xml": "<a>" * 600 + "</a>" * 600,
        },
    )
    for path in (sheet_bomb, tree_bomb, deep):
        with pytest.raises(ValueError, match = "too large to index|nested too deeply"):
            parsers.parse(str(path))


@pytest.mark.parametrize("format_id", [14, 31, 57, 75])
def test_xlsx_builtin_date_formats_include_locale_ids(tmp_path, format_id):
    serial = (dt.date(2026, 3, 31) - dt.date(1899, 12, 30)).days
    path = _zip(
        tmp_path / "dates.xlsx",
        {
            "xl/workbook.xml": f'<workbook {S}><sheets><sheet name="D" sheetId="1" r:id="rId1"/></sheets></workbook>',
            "xl/_rels/workbook.xml.rels": f'<Relationships {REL}><Relationship Id="rId1" Target="worksheets/sheet1.xml"/></Relationships>',
            "xl/styles.xml": f'<styleSheet {S}><cellXfs><xf numFmtId="0"/><xf numFmtId="{format_id}"/></cellXfs></styleSheet>',
            "xl/worksheets/sheet1.xml": f'<worksheet {S}><sheetData><row r="1"><c r="A1" s="1"><v>{serial}</v></c></row></sheetData></worksheet>',
        },
    )
    assert _text(path) == "Sheet: D\n2026-03-31"


def test_pptx_reads_chart_titles_and_cached_data(tmp_path):
    C = 'xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
    pt = lambda i, v: f'<c:pt idx="{i}"><c:v>{v}</c:v></c:pt>'
    series = lambda name, values: (
        f"<c:ser><c:tx><c:strRef><c:strCache>{pt(0, name)}</c:strCache></c:strRef></c:tx>"
        f"<c:cat><c:strRef><c:strCache>{pt(0, 'North')}{pt(1, 'South')}</c:strCache></c:strRef></c:cat>"
        f"<c:val><c:numRef><c:numCache>{pt(0, values[0])}{pt(1, values[1])}</c:numCache></c:numRef></c:val></c:ser>"
    )
    path = _zip(
        tmp_path / "chart.pptx",
        {
            "ppt/presentation.xml": f'<p:presentation {P}><p:sldIdLst><p:sldId id="256" r:id="rId2"/></p:sldIdLst></p:presentation>',
            "ppt/_rels/presentation.xml.rels": f'<Relationships {REL}><Relationship Id="rId2" Target="slides/slide1.xml"/></Relationships>',
            "ppt/slides/slide1.xml": f"<p:sld {P}><p:cSld><p:spTree><p:graphicFrame><a:graphic><a:graphicData>"
            '<c:chart xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" r:id="rId3"/>'
            "</a:graphicData></a:graphic></p:graphicFrame></p:spTree></p:cSld></p:sld>",
            "ppt/slides/_rels/slide1.xml.rels": f'<Relationships {REL}><Relationship Id="rId3" Target="../charts/chart1.xml"/></Relationships>',
            "ppt/charts/chart1.xml": f"<c:chartSpace {C}><c:chart><c:title><c:tx><c:rich><a:p><a:r><a:t>Sales zebramarker</a:t></a:r></a:p></c:rich></c:tx></c:title>"
            f"<c:plotArea><c:barChart>{series('Q1', (1200, 900))}{series('Q2', (1300, 950))}</c:barChart></c:plotArea></c:chart></c:chartSpace>",
        },
    )
    assert _text(path) == "Sales zebramarker\nQ1 | Q2\nNorth | 1200 | 1300\nSouth | 900 | 950"


def test_odp_reads_one_page_per_slide(tmp_path):
    pages = parsers.parse(str(build_odp(tmp_path / "slides.odp")))
    assert [(p.page_number, p.text) for p in pages] == [
        (1, "First slide\nNotes:\nzebramarker notes"),
        (2, "Second slide revenue"),
    ]


def test_epub_follows_the_spine_and_skips_navigation(tmp_path):
    text = _text(build_epub(tmp_path / "book.epub"))
    assert "navmarker" not in text
    assert text.index("Revenue increased") < text.index("Zebramarker ends the book")


@pytest.mark.parametrize(
    "uri, protected", [("OEBPS/text/Chapter%201.xhtml", True), ("OEBPS/fonts/a.otf", False)]
)
def test_epub_with_encrypted_chapters_is_refused(tmp_path, uri, protected):
    book = build_epub(tmp_path / "book.epub")
    with zipfile.ZipFile(book, "a") as z:
        z.writestr(
            "META-INF/encryption.xml",
            '<encryption xmlns="urn:oasis:names:tc:opendocument:xmlns:container" '
            'xmlns:enc="http://www.w3.org/2001/04/xmlenc#"><enc:EncryptedData>'
            f'<enc:CipherData><enc:CipherReference URI="{uri}"/></enc:CipherData>'
            "</enc:EncryptedData></encryption>",
        )
    if protected:
        with pytest.raises(ValueError, match = "DRM protected"):
            parsers.parse(str(book))
    else:  # obfuscated fonts are listed too, and never block reading
        assert "Zebramarker" in _text(book)


def test_eml_decodes_headers_and_body(tmp_path):
    text = _text(build_eml(tmp_path / "mail.eml"))
    assert "Subject: Q1 numbers résumé" in text
    assert "From: Ann <ann@example.com>" in text
    assert "Attachments: data.csv" in text
    assert "Zebramarker in the body." in text
    assert "html copy" not in text


@pytest.mark.parametrize("extension", [".docx", ".docm", ".dotx", ".dotm"])
def test_word_macro_and_template_files_read_like_docx(tmp_path, extension):
    text = _text(build_word(tmp_path / f"report{extension}"))
    assert "Quarterly report" in text and "Zebramarker" in text


def test_word_extension_on_another_package_is_refused(tmp_path):
    workbook = build_word(tmp_path / "report.docx").read_bytes()
    with zipfile.ZipFile(io.BytesIO(workbook)) as source:
        types = source.read("[Content_Types].xml").replace(
            _WORD_TYPES[".docx"].encode(),
            b"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
        )
        with zipfile.ZipFile(tmp_path / "sheet.docm", "w") as target:
            for info in source.infolist():
                target.writestr(
                    info, types if info.filename == "[Content_Types].xml" else source.read(info)
                )
    with pytest.raises(ValueError, match = "not a Word file"):
        parsers.parse(str(tmp_path / "sheet.docm"))


def test_mhtml_reads_the_page_and_skips_resources(tmp_path):
    text = _text(build_mhtml(tmp_path / "page.mhtml"))
    assert "Zebramarker on the page." in text
    assert "PNG" not in text


def test_xhtml_reads_visible_text(tmp_path):
    text = _text(build_xhtml(tmp_path / "page.xhtml"))
    assert "Zebramarker résumé" in text and "p{}" not in text


def test_rtf_keeps_text_and_drops_control_groups(tmp_path):
    assert _text(build_rtf(tmp_path / "memo.rtf")) == (
        "Quarterly report\nCafé € revenue increased\nlink text\nZebramarker | 1200"
    )


def test_doc_reads_compressed_and_unicode_pieces_and_note_stories(tmp_path):
    assert _text(build_doc(tmp_path / "memo.doc")) == (
        "Quarterly report\nCafé 東京 link Zebramarker | 1200\n\nBoxed text"
        "\n\nFootnotes:\nFootnote detail\n\nEndnotes:\nEndnote source"
    )


def test_xls_reads_errors_and_continued_formula_strings(tmp_path):
    long_text = "A" * 8000 + "Ж" * 300
    string = struct.pack("<HB", len(long_text), 0) + b"A" * 8000
    sheet = (
        _record(0x0809, b"\0" * 16)
        + _record(0x0205, struct.pack("<HHHBB", 0, 0, 0, 0x07, 1))  # BOOLERR #DIV/0!
        + _record(0x0006, struct.pack("<HHH", 0, 1, 0) + b"\x02\0\x2a\0\0\0\xff\xff" + b"\0" * 6)
        + _record(0x0006, struct.pack("<HHH", 1, 0, 0) + b"\0" * 6 + b"\xff\xff" + b"\0" * 6)
        + _record(0x0207, string)
        + _record(0x003C, b"\x01" + "Ж".encode("utf-16-le") * 300)
        + _record(0x000A, b"")
    )
    globals_ = (
        _record(0x0809, b"\0" * 16)
        + _record(0x0085, b"\0" * 4 + b"\0\0" + struct.pack("<B", 1) + b"\0S")
        + _record(0x000A, b"")
    )
    stream = bytearray(globals_ + sheet)
    struct.pack_into("<I", stream, 4 + 16 + 4, len(globals_))
    path = tmp_path / "errors.xls"
    path.write_bytes(compound_file({("Workbook",): bytes(stream)}))
    assert _text(path) == f"Sheet: S\n#DIV/0! | #N/A\n{long_text}"


def test_xls_reads_continued_strings_numbers_formulas_and_dates(tmp_path):
    assert _text(build_xls(tmp_path / "book.xls")) == (
        "Sheet: Revenue\nRegion | 1200.5\nZebramarker€ | -7\n3 | 4\nformula text | TRUE"
        "\n2026-03-31 | 2026-04-01"
    )


def test_xls_reads_excel_95_byte_strings(tmp_path):
    def text(value, size):
        raw = value.encode("cp1251")
        return struct.pack("<B" if size == 1 else "<H", len(raw)) + raw

    bof = _record(0x0809, struct.pack("<HH", 0x0500, 0x0005) + b"\0" * 4)
    globals_ = lambda offset: (
        bof
        + _record(0x0042, struct.pack("<H", 1251))  # CODEPAGE
        + _record(0x0085, struct.pack("<IBB", offset, 0, 0) + text("Выручка", 1))
        + _record(0x000A, b"")
    )
    sheet = (
        bof
        + _record(0x0204, struct.pack("<HHH", 0, 0, 0) + text("Zebramarker", 2))
        + _record(0x0204, struct.pack("<HHH", 0, 1, 0) + text("Привет", 2))
        + _record(0x0006, struct.pack("<HHH", 1, 0, 0) + b"\0" * 6 + b"\xff\xff" + b"\0" * 6)
        + _record(0x0207, text("formula text", 2))
        + _record(0x000A, b"")
    )
    path = tmp_path / "excel95.xls"
    path.write_bytes(compound_file({("Book",): globals_(len(globals_(0))) + sheet}))
    assert _text(path) == "Sheet: Выручка\nZebramarker | Привет\nformula text"


def test_ppt_reads_slide_text_and_skips_masters(tmp_path):
    pages = parsers.parse(str(build_ppt(tmp_path / "deck.ppt")))
    assert [(p.page_number, p.text) for p in pages] == [
        (1, "Quarterly review\nRevenue up"),
        (2, "Zebramarker slide"),
    ]


@pytest.mark.parametrize("broken_link", [False, True])
def test_ppt_reads_the_latest_save_with_text_boxes_and_notes(tmp_path, broken_link):
    pages = parsers.parse(str(build_saved_ppt(tmp_path / "saved.ppt", broken_link = broken_link)))
    assert [(p.page_number, p.text) for p in pages] == [
        (1, "Title from list\nText box zebramarker\nNotes:\nSpeaker notes")
    ]


def test_ppt_falls_back_to_slide_drawings(tmp_path):
    pages = parsers.parse(str(build_ppt(tmp_path / "deck.ppt", from_slides = True)))
    assert [p.text for p in pages] == ["Drawn slide", "Zebramarker drawn"]


def test_msg_reads_headers_body_and_attachment_names(tmp_path):
    assert _text(build_msg(tmp_path / "mail.msg")) == (
        "From: Ann <ann@example.com>\nTo: Bob\nDate: 2026-10-06 10:00 UTC\nSubject: Q1 numbers\n"
        "Attachments: data.csv\n\nRevenue increased.\nZebramarker in the body."
    )


def test_msg_uses_the_smtp_address_for_exchange_senders(tmp_path):
    utf16 = lambda text: text.encode("utf-16-le")
    path = tmp_path / "exchange.msg"
    path.write_bytes(
        compound_file(
            {
                ("__substg1.0_0C1A001F",): utf16("Ann"),
                ("__substg1.0_0C1F001F",): utf16("/O=EXAMPLE/OU=EXCHANGE/CN=RECIPIENTS/CN=ANN"),
                ("__substg1.0_5D01001F",): utf16("ann@example.com"),
                ("__substg1.0_1000001F",): utf16("Body zebramarker"),
            }
        )
    )
    assert _text(path) == "From: Ann <ann@example.com>\n\nBody zebramarker"


def test_eml_reads_every_inline_part(tmp_path):
    message = EmailMessage()
    message["Subject"] = "Two parts"
    message.set_content("first part")
    message.make_mixed()
    inline = EmailMessage()
    inline.set_content("second zebramarker")
    inline["Content-Disposition"] = "inline"
    message.attach(inline)
    message.add_attachment(b"x,y\n", maintype = "text", subtype = "csv", filename = "data.csv")
    path = tmp_path / "mixed.eml"
    path.write_bytes(bytes(message))
    text = _text(path)
    assert text.endswith("first part\n\nsecond zebramarker")
    assert "x,y" not in text


def test_eml_falls_back_to_html_when_plain_is_empty(tmp_path):
    message = EmailMessage()
    message["Subject"] = "HTML only"
    message.set_content("   \n")
    message.add_alternative("<p>zebramarker in html</p>", subtype = "html")
    path = tmp_path / "html.eml"
    path.write_bytes(bytes(message))
    assert _text(path).endswith("zebramarker in html")


def test_msg_decodes_ansi_properties_and_compressed_rtf_body(tmp_path):
    assert _text(build_ansi_msg(tmp_path / "ansi.msg")) == (
        "From: Иван\nSubject: Привет\n\nZebramarker in the RTF body\nrepeat repeat repeat"
    )


def _strict(source, target):
    """Rewrite a Transitional OOXML package with the Strict namespaces."""
    pairs = [
        (
            b"http://schemas.openxmlformats.org/spreadsheetml/2006/main",
            b"http://purl.oclc.org/ooxml/spreadsheetml/main",
        ),
        (
            b"http://schemas.openxmlformats.org/presentationml/2006/main",
            b"http://purl.oclc.org/ooxml/presentationml/main",
        ),
        (
            b"http://schemas.openxmlformats.org/drawingml/2006/main",
            b"http://purl.oclc.org/ooxml/drawingml/main",
        ),
        (
            b"http://schemas.openxmlformats.org/officeDocument/2006/relationships",
            b"http://purl.oclc.org/ooxml/officeDocument/relationships",
        ),
    ]
    with zipfile.ZipFile(source) as z:
        members = {name: z.read(name) for name in z.namelist()}
    for name, data in members.items():
        for old, new in pairs:
            data = data.replace(old, new)
        members[name] = data
    return _zip(target, members)


@pytest.mark.parametrize("builder, extension", [(build_xlsx, ".xlsx"), (build_pptx, ".pptx")])
def test_strict_ooxml_reads_like_transitional(tmp_path, builder, extension):
    transitional = builder(tmp_path / f"transitional{extension}")
    strict = _strict(transitional, tmp_path / f"strict{extension}")
    with zipfile.ZipFile(strict) as z:
        assert all(b"2006/main" not in z.read(name) for name in z.namelist())
    assert _text(strict) == _text(transitional) != ""


def test_odt_skips_tracked_deletions_and_reads_nested_tables_once(tmp_path):
    path = _odf(
        tmp_path / "changes.odt",
        "<office:text><text:tracked-changes><text:changed-region>"
        "<text:deletion><text:p>deletedmarker</text:p></text:deletion>"
        "</text:changed-region></text:tracked-changes>"
        '<text:p>See <text:bookmark-ref text:ref-name="b">Chapter 2</text:bookmark-ref>.</text:p>'
        "<table:table><table:table-header-rows><table:table-row>"
        "<table:table-cell><text:p>Head</text:p></table:table-cell>"
        "</table:table-row></table:table-header-rows><table:table-row><table:table-cell>"
        "<table:table><table:table-row><table:table-cell><text:p>inner</text:p></table:table-cell>"
        "<table:table-cell><text:p>cell</text:p></table:table-cell></table:table-row></table:table>"
        "</table:table-cell><table:table-cell><text:p>outer</text:p></table:table-cell>"
        "</table:table-row></table:table></office:text>",
    )
    assert _text(path) == "See Chapter 2.\nHead\ninner | cell | outer"


def test_rtf_decodes_literal_bytes_in_the_document_code_page(tmp_path):
    path = tmp_path / "ru.rtf"
    path.write_bytes(b"{\\rtf1\\ansi\\ansicpg1251 " + "Привет".encode("cp1251") + b" \\'80}")
    assert _text(path) == "Привет Ђ"


def test_compound_file_reads_mini_and_regular_streams():
    small, large = b"a" * 100, bytes(range(256)) * 40
    reader = cfb.CompoundFile(compound_file({("small",): small, ("dir", "large"): large}))
    assert reader.open("small") == small
    assert reader.open("DIR", "Large") == large
    assert sorted(reader.listdir()) == ["dir", "small"]


def test_compound_file_reads_each_fat_sector_once_per_file_sector():
    # Every DIFAT slot names the same FAT sector: the FAT must stay the size of the file.
    data = bytearray(compound_file({("small",): b"a" * 100}))
    n_difat = 64
    first = len(data) // 512 - 1
    data += b"".join(
        struct.pack("<128I", *([0] * 127 + [first + i + 1 if i + 1 < n_difat else 0xFFFFFFFE]))
        for i in range(n_difat)
    )
    struct.pack_into("<I", data, 0x2C, 0)  # no FAT sector count to trim by
    struct.pack_into("<II", data, 0x44, first, n_difat)
    reader = cfb.CompoundFile(bytes(data))
    assert len(reader._fat) <= len(data) // 4
    assert reader.open("small") == b"a" * 100


def test_encrypted_doc_is_refused(tmp_path):
    with pytest.raises(ValueError, match = "password"):
        parsers.parse(str(build_doc(tmp_path / "secret.doc", encrypted = True)))


@pytest.mark.parametrize("extension", [".doc", ".xls", ".ppt", ".msg"])
def test_non_compound_file_is_refused(tmp_path, extension):
    path = tmp_path / f"fake{extension}"
    path.write_bytes(b"not a compound file" * 50)
    with pytest.raises(ValueError):
        parsers.parse(str(path))


@pytest.mark.parametrize("extension", [".xlsx", ".pptx", ".odt", ".epub"])
def test_damaged_archive_is_refused(tmp_path, extension):
    path = tmp_path / f"broken{extension}"
    path.write_bytes(b"PK\x03\x04 truncated")
    with pytest.raises(ValueError):
        parsers.parse(str(path))


def test_rtf_pairs_surrogates_and_skips_tracked_deletions(tmp_path):
    path = tmp_path / "changes.rtf"
    # As macOS textutil writes an emoji; the deletion as Word writes a tracked change.
    path.write_bytes(
        rb"{\rtf1\ansi Party \uc0\u55356 \u57225  kept{\deleted\revauthdel1 removedmarker \u8364?}"
        rb" end\deleted gone\deleted0  back\plain}"
    )
    assert _text(path) == "Party \U0001f389 kept end back"


def test_compressed_rtf_stops_at_the_declared_size():
    # One literal, then back-references that would expand to ~5 MB.
    body = b"\x00A" + (b"\xff" + b"\x00\x0f" * 8) * 40_000
    data = struct.pack("<IIII", len(body) + 12, 1000, 0x75465A4C, 0) + body
    assert len(office_formats._decompress_rtf(data)) == 1000
    with pytest.raises(ValueError, match = "too large"):
        office_formats._decompress_rtf(struct.pack("<IIII", 16, 2**31, 0x75465A4C, 0))


def test_rtf_decodes_bytes_with_the_font_charset(tmp_path):
    path = tmp_path / "fonts.rtf"
    path.write_bytes(
        rb"{\rtf1\ansi\ansicpg1252\deff0{\fonttbl{\f0\fswiss\fcharset0 Arial;}"
        rb"{\f1\froman\fcharset204 Times New Roman Cyr;}{\f2\fnil\cpg1253 Greek;}}"
        rb"\f0 Hello {\f1 \'cf\'f0\'e8\'e2\'e5\'f2} back {\f2 \'e1} \plain\'e9\par}"
    )
    assert _text(path) == "Hello Привет back α é"


def test_rtf_prefers_the_unicode_branch_of_upr(tmp_path):
    path = tmp_path / "upr.rtf"
    path.write_bytes(rb"{\rtf1\ansi \upr{?}{\*\ud{\u1040?}} end}")
    assert _text(path) == "\u0410 end"


def test_rtf_refuses_runaway_nesting(tmp_path):
    path = tmp_path / "deep.rtf"
    path.write_bytes(b"{\\rtf1 " + b"{" * 5000)
    with pytest.raises(ValueError, match = "nested too deeply"):
        parsers.parse(str(path))


def test_rtf_negative_binary_length_is_ignored(tmp_path):
    path = tmp_path / "bin.rtf"
    path.write_bytes(rb"{\rtf1\ansi before \bin-999999999 after}")
    assert _text(path) == "before after"


def test_ods_refuses_runaway_repeats(tmp_path):
    cell = "x" * 100_000
    path = _odf(
        tmp_path / "repeats.ods",
        '<office:spreadsheet><table:table table:name="S">'
        '<table:table-row table:number-rows-repeated="1000">'
        f'<table:table-cell table:number-columns-repeated="1000"><text:p>{cell}</text:p></table:table-cell>'
        "</table:table-row></table:table></office:spreadsheet>",
    )
    with pytest.raises(ValueError, match = "repeats too much"):
        parsers.parse(str(path))


def test_xml_entities_are_refused(tmp_path):
    path = _zip(
        tmp_path / "bomb.odt", {"content.xml": '<!DOCTYPE x [<!ENTITY a "aaaa">]><x>&a;</x>'}
    )
    with pytest.raises(ValueError):
        parsers.parse(str(path))


@pytest.mark.parametrize("extension", sorted(BUILDERS))
def test_documents_upload_index_and_stay_searchable(rag_home, stub_embeddings, tmp_path, extension):
    if not rag_db.rag_available():  # vec0 will not load on some macOS Pythons
        pytest.skip("sqlite-vec unavailable here")
    source = BUILDERS[extension](tmp_path / f"report{extension}")
    stored, filename, _hash = _save_upload(
        UploadFile(file = io.BytesIO(source.read_bytes()), filename = source.name)
    )
    assert filename == source.name
    scope = store.thread_scope("office")
    doc_id, job = ingestion.start_ingestion(
        scope, None, "office", filename, stored, background = False
    )
    assert ingestion.get_job_status(job)["status"] == "completed"
    conn = rag_db.get_connection()
    try:
        assert store.get_document(conn, doc_id)["status"] == "completed"
        assert store.search_lexical(conn, scope, "zebramarker", 5)
    finally:
        conn.close()
