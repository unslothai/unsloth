# SPDX-License-Identifier: AGPL-3.0-only
# Copyright 2026-present the Unsloth AI Inc. team. All rights reserved. See /studio/LICENSE.AGPL-3.0

from __future__ import annotations

import itertools
import sys
import threading
from pathlib import Path

import pytest

_BACKEND_DIR = str(Path(__file__).resolve().parent.parent)
if _BACKEND_DIR not in sys.path:
    sys.path.insert(0, _BACKEND_DIR)

pytest.importorskip("docx")
import docx  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

from core.rag import parsers  # noqa: E402


def _child(tag, **attrs):
    element = OxmlElement(tag)
    for key, value in attrs.items():
        element.set(qn(key), value)
    return element


def _children():
    text = _child("w:t")
    text.text = "word"
    spaced = _child("w:t")
    spaced.text = " gap "
    deleted = _child("w:delText")
    deleted.text = "gone"
    instr = _child("w:instrText")
    instr.text = " PAGE "
    return [
        text,
        spaced,
        _child("w:tab"),
        _child("w:br"),
        _child("w:br", **{"w:type": "page"}),
        _child("w:br", **{"w:type": "column"}),
        _child("w:br", **{"w:type": "textWrapping"}),
        _child("w:cr"),
        _child("w:noBreakHyphen"),
        _child("w:softHyphen"),
        _child("w:ptab", **{"w:relativeTo": "margin", "w:alignment": "right", "w:leader": "none"}),
        deleted,
        instr,
    ]


def _run_with(parts):
    document = docx.Document()
    run = document.add_paragraph().add_run()
    for part in parts:
        run._r.append(part)
    return run


def test_run_text_matches_python_docx_for_every_text_child():
    """Without checkbox fields or symbols, the run text is exactly python-docx's Run.text."""
    kinds = len(_children())
    for i, j in itertools.product(range(kinds), repeat = 2):
        fresh = _children()
        run = _run_with([fresh[i], fresh[j]])
        assert parsers._docx_run_text(run._r) == run.text, (i, j)
    run = _run_with(_children())
    assert parsers._docx_run_text(run._r) == run.text


def test_run_text_is_the_same_on_every_thread():
    run = _run_with(_children())
    expected = run.text
    results = []

    def read():
        results.extend(parsers._docx_run_text(run._r) for _ in range(50))

    threads = [threading.Thread(target = read) for _ in range(4)]
    for thread in threads:
        thread.start()
    for thread in threads:
        thread.join()
    assert results == [expected] * 200
