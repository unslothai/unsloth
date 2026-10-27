# SPDX-License-Identifier: AGPL-3.0-only
# Copyright 2026-present the Unsloth AI Inc. team. All rights reserved. See /studio/LICENSE.AGPL-3.0

"""Document parsing -> list[Page], one dispatch with lazy optional deps.

PDFs keep per-page boundaries (``page_number``); txt/md/docx/html return a single
page. ``parse(path, want_images=True)`` also returns embedded images. Heavy imports
are lazy, so importing this module never fails on a missing dep.
"""

from __future__ import annotations

import codecs
import logging
import os
import re
import secrets
import struct
from dataclasses import dataclass
from html.parser import HTMLParser

from . import config

logger = logging.getLogger(__name__)


@dataclass(frozen = True)
class Page:
    """A unit of extracted text. ``page_number`` is 1-based (None if N/A)."""

    text: str
    page_number: int | None = None
    char_count: int = 0
    needs_ocr: bool = False


@dataclass(frozen = True)
class ParsedImage:
    """A raster image embedded in a document (PDF only)."""

    image_bytes: bytes
    page_number: int | None
    xref: int
    full_page: bool = False
    tile_index: int | None = None
    tile_count: int = 0


def _page(text: str, page_number: int | None) -> Page:
    return Page(text = text, page_number = page_number, char_count = len(text))


_HTML_SKIP_TAGS = frozenset(("script", "style", "template"))
_HTML_BLOCK_TAGS = frozenset(
    "address article aside blockquote br caption center dd details dialog dir div dl dt"
    " fieldset figcaption figure footer form h1 h2 h3 h4 h5 h6 header hgroup hr legend li"
    " listing main menu nav ol optgroup option p plaintext pre search section summary table"
    " td text textarea th title tr ul xmp".split()
)
_HTML_PRE_TAGS = frozenset(("listing", "plaintext", "pre", "textarea", "xmp"))
# atomic inline boxes separate neighbouring text without breaking the line
_HTML_BOX_TAGS = frozenset(("button", "img", "input", "select"))
_HTML_ROW_GROUPS = frozenset(("thead", "tbody", "tfoot"))
_HTML_LIST_TAGS = frozenset(("dir", "menu", "ol", "ul"))
_HTML_LIST_FORMATS = {
    "1": "decimal",
    "a": "lowerAlpha",
    "A": "upperAlpha",
    "i": "lowerRoman",
    "I": "upperRoman",
}
_HTML_TABLE_DEPTH = 32


class _HtmlPre(str):
    pass


class _HtmlTable:
    def __init__(
        self,
        sink: list,
        nested: bool = False,
    ) -> None:
        self.sink = sink
        self.nested = nested
        self.row: list[list | None] | None = None
        self.spans: list[int] = []


def _html_int(value) -> int | None:
    match = re.match(r"\s*([+-]?\d{1,18})(?!\d)", value or "")
    return int(match.group(1)) if match else None


def _html_span(attrs, name: str, limit: int) -> int:
    try:
        value = int(dict(attrs).get(name) or 1)
    except ValueError:
        return 1
    # rowspan="0" runs to the end of its row group
    return limit if value == 0 and name == "rowspan" else min(max(value, 1), limit)


class _Stripper(HTMLParser):
    """collect visible text, one line per block element."""

    def __init__(self, span_budget: int = 0) -> None:
        super().__init__()
        self._skip = 0
        self._pre = 0
        self._templates: list[bool] = []
        self._line: list[str] = []
        self.out: list[str] = []
        self._tables: list[_HtmlTable] = []
        self._span_budget = span_budget
        self._deep = 0
        self._lists: list[tuple[str, dict | None, list]] = []
        self._open_lists: dict[str, int] = {}
        self.labels: list[str] = []
        self.nonce = secrets.token_hex(4)
        self._label: str | None = None

    def _in_table(self) -> bool:
        return bool(self._tables) and not self._skip and not self._deep

    def _end_cell(self) -> None:
        if self.out is not self._tables[-1].sink:
            self._emit_label()
        self._flush()
        self.out = self._tables[-1].sink

    def _end_row(self) -> None:
        self._end_cell()
        table = self._tables[-1]
        row, table.row = table.row, None
        if row is None:
            return
        missing = max(len(table.spans) - len(row), 0)
        for col in range(len(row), len(table.spans)):
            table.spans[col] = max(table.spans[col] - 1, 0)
        self._span_budget -= missing
        if self._span_budget < 0:
            table.spans.clear()
        else:
            row.extend([None] * missing)
        while table.spans and not table.spans[-1]:
            table.spans.pop()
        if len(row) == 1:
            table.sink.extend(line for line in row[0] or () if line is not None)
            return
        cells: list[str] = []
        below: list[str] = []
        for cell in row:
            lines = cell or []
            parent: list[str] = []
            in_nested = False
            for line in lines:
                if line is None:
                    in_nested = not in_nested
                elif in_nested:
                    below.append(line)
                else:
                    parent.append(line)
            parts: list[str] = []
            previous_pre = False
            for line in parent:
                is_pre = isinstance(line, _HtmlPre)
                if parts:
                    parts.append("\n" if previous_pre or is_pre else " ")
                parts.append(str(line) if is_pre else " ".join(line.split()))
                previous_pre = is_pre
            cells.append("".join(parts))
        if any(cells):
            table.sink.append(" | ".join(cells))
        table.sink.extend(below)

    def _start_cell(self, attrs) -> None:
        self._end_cell()
        table = self._tables[-1]
        if table.row is None:
            table.row = []
        row, spans = table.row, table.spans
        start = len(row)
        while len(row) < len(spans) and spans[len(row)]:
            spans[len(row)] -= 1
            row.append(None)
        self._flush()
        cell: list[str] = []
        rowspan = _html_span(attrs, "rowspan", 65534)
        colspan = _html_span(attrs, "colspan", 1000)
        # each spanned slot adds an empty field, so stop when fields outnumber input characters.
        self._span_budget -= len(row) - start + colspan - 1
        if self._span_budget < 0:
            rowspan = colspan = 1
            spans[:] = [0] * len(spans)
        for i in range(colspan):
            if len(row) < len(spans):
                spans[len(row)] = rowspan - 1
            else:
                spans.append(rowspan - 1)
            row.append(None if i else cell)
        self.out = cell

    def _end_list(self) -> str:
        tag, attrs, items = self._lists.pop()
        self._open_lists[tag] -= 1
        if attrs is not None:
            backwards = "reversed" in attrs
            number = _html_int(attrs.get("start"))
            if number is None:
                number = len(items) if backwards else 1
            kind = _HTML_LIST_FORMATS.get(attrs.get("type") or "")
            for index, value, item_kind in items:
                number = number if value is None else value
                item_format = _HTML_LIST_FORMATS.get(item_kind or "", kind)
                self.labels[index] = _list_number(number, item_format) + "."
                number += -1 if backwards else 1
        return tag

    def _emit_label(self) -> None:
        if self._label:
            self._flush()
            self._line.append(self._label)
            self._label = None
            self._flush()

    def _flush(self) -> None:
        text = "".join(self._line)
        self._line = []
        # whitespace inside <pre>/<textarea> is content; elsewhere it is layout
        text = text.strip("\n") if self._pre else " ".join(text.split())
        if text.strip():
            self.out.append(_HtmlPre(text) if self._pre else text)

    def handle_starttag(self, tag, attrs):
        if tag == "template":
            # declarative shadow roots render; other templates are inert
            inert = (dict(attrs).get("shadowrootmode") or "").lower() not in ("open", "closed")
            self._templates.append(inert)
            self._skip += inert
        elif tag in _HTML_SKIP_TAGS:
            self._skip += 1
        elif tag in ("td", "th") and self._in_table():
            self._start_cell(attrs)
        elif tag == "tr" and self._in_table():
            self._end_row()
            self._flush()
            self._tables[-1].row = []
        elif tag in _HTML_ROW_GROUPS and self._in_table():
            self._end_row()
            self._tables[-1].spans = []
        elif tag in _HTML_BLOCK_TAGS and not self._skip:
            self._flush()
            # each nesting level copies lines into its parent, so deeply nested tables read as plain blocks
            if tag == "table" and (self._deep or len(self._tables) >= _HTML_TABLE_DEPTH):
                self._deep += 1
            elif tag == "table":
                nested = bool(self._tables and self.out is not self._tables[-1].sink)
                if nested:
                    self.out.append(None)
                self._tables.append(_HtmlTable(self.out, nested))
            if tag in _HTML_PRE_TAGS:
                self._pre += 1
            if tag in _HTML_LIST_TAGS or tag == "li":
                self._emit_label()
            if tag in _HTML_LIST_TAGS:
                self._lists.append((tag, dict(attrs) if tag == "ol" else None, []))
                self._open_lists[tag] = self._open_lists.get(tag, 0) + 1
            elif tag == "li" and self._lists and self._lists[-1][1] is not None:
                attrs = dict(attrs)
                item = (len(self.labels), _html_int(attrs.get("value")), attrs.get("type"))
                self._lists[-1][2].append(item)
                self._label = f"\ue000{self.nonce}.{len(self.labels)}\ue001 "
                self.labels.append("")
        elif tag in _HTML_BOX_TAGS and not self._skip and not self._pre:
            self._line.append(" ")
        elif tag == "tspan" and not self._skip and any(k in ("x", "y") for k, _ in attrs):
            # absolute x/y coordinates start a separate SVG text label or line
            self._flush()

    def handle_endtag(self, tag):
        if tag == "template":
            if self._templates and self._templates.pop() and self._skip:
                self._skip -= 1
        elif tag in _HTML_SKIP_TAGS:
            if self._skip:
                self._skip -= 1
        elif tag in ("td", "th") and self._in_table():
            self._end_cell()
        elif tag in _HTML_ROW_GROUPS and self._in_table():
            self._end_row()
            self._tables[-1].spans = []
        elif tag in ("tr", "table") and self._in_table():
            self._end_row()
            if tag == "table":
                table = self._tables.pop()
                if table.nested:
                    self.out.append(None)
            self._flush()
        elif tag in _HTML_BLOCK_TAGS and not self._skip:
            self._flush()
            if tag == "table" and self._deep:
                self._deep -= 1
            if tag in _HTML_PRE_TAGS and self._pre:
                self._pre -= 1
            if self._open_lists.get(tag):
                self._emit_label()
                while self._end_list() != tag:
                    pass
        elif tag in _HTML_BOX_TAGS and not self._skip and not self._pre:
            self._line.append(" ")

    def handle_data(self, data):
        if not self._skip:
            if self._label and data.strip():
                self._line.append(self._label)
                self._label = None
            self._line.append(data)

    def close(self):
        super().close()
        while self._tables:
            self._end_row()
            table = self._tables.pop()
            if table.nested:
                self.out.append(None)
        self._emit_label()
        self._flush()
        while self._lists:
            self._end_list()


def _html(raw: str) -> list[Page]:
    parser = _Stripper(len(raw))
    parser.feed(raw)
    parser.close()
    label = re.compile(f"\ue000{parser.nonce}\\.(\\d+)\ue001")
    text = label.sub(lambda match: parser.labels[int(match.group(1))], "\n".join(parser.out))
    return [_page(text, 1)]


# pymupdf4llm can corrupt positioned RTL and Indic glyphs; use logical-order
# get_text() at unslothai/unsloth#5351 thresholds
_SHAPED_PRESENTATION_FORMS = re.compile("[\ufb1d-\ufdff\ufe70-\ufefc]")
_PDF_FALLBACK_MIN_BAD_GLYPHS = 5
_PDF_FALLBACK_BAD_GLYPH_RATIO = 0.0005
_PDF_INCOMPLETE_RATIO = 0.75
_PDF_INCOMPLETE_MIN_LETTERS = 200


def _markdown_corrupted(text: str) -> bool:
    """True when pymupdf4llm's glyph reconstruction mangled the text: shaped RTL
    Presentation Forms or U+FFFD replacements above a small floor/ratio (so a lone
    legitimate shaped glyph does not force the fallback)."""
    if not text:
        return False
    threshold = max(_PDF_FALLBACK_MIN_BAD_GLYPHS, _PDF_FALLBACK_BAD_GLYPH_RATIO * len(text))
    shaped = len(_SHAPED_PRESENTATION_FORMS.findall(text))
    return shaped > threshold or text.count("\ufffd") > threshold


def _markdown_incomplete(markdown: str, plain: str) -> bool:
    """True when ``markdown`` holds far fewer letters than the raw ``get_text`` layer -- a
    coarse guard for heavy-RTL pages pymupdf4llm silently drops without shaped glyphs."""
    plain_letters = sum(1 for c in plain if c.isalnum())
    if plain_letters < _PDF_INCOMPLETE_MIN_LETTERS:
        return False
    markdown_letters = sum(1 for c in markdown if c.isalnum())
    return markdown_letters < _PDF_INCOMPLETE_RATIO * plain_letters


# JPEG2000 decodes at ~45 ms per megapixel; a logo stays far below this, a 300 dpi scan far above.
_SCAN_IMAGE_MIN_PIXELS = 1_000_000


def _image_covers_a_corner(doc, number: int) -> bool:
    """True when a scan-sized image covers a corner pymupdf4llm samples, since that render holds the GIL."""
    try:
        import fitz

        page = doc[number]
        r = page.rect
        # 1pt slack: the render clip rounds outward to whole pixels.
        corners = [
            fitz.Rect(x - 1, y - 1, x + 11, y + 11)
            for x in (r.x0, r.x1 - 10)
            for y in (r.y0, r.y1 - 10)
        ]
        # Image boxes are unrotated; the probe clips page.rect, which is rotated.
        boxes = [
            fitz.Rect(info["bbox"]) * page.rotation_matrix
            for info in page.get_image_info()
            if info["width"] * info["height"] >= _SCAN_IMAGE_MIN_PIXELS
        ]
    except Exception:
        return False
    return any(box.intersects(corner) for box in boxes for corner in corners)


def _pdf_markdown(doc, pages: range | None = None) -> list[str] | None:
    """Per-page layout-aware Markdown (tables, headings, lists) via pymupdf4llm; index
    i maps to page i+1. Returns None when the lib is missing, extraction fails, or the
    page count does not line up, so the caller falls back to plain PyMuPDF text."""
    try:
        import pymupdf4llm
    except Exception:
        return None
    try:
        kwargs = {"page_chunks": True, "show_progress": False, "ignore_images": True}
        indices = list(range(doc.page_count) if pages is None else pages)
        scans = {i for i in indices if _image_covers_a_corner(doc, i)}
        if not scans:
            if pages is not None:
                kwargs["pages"] = indices
            chunks = pymupdf4llm.to_markdown(doc, **kwargs)
        else:
            # Background detection only where it is cheap. Each call still builds its header
            # table from the whole (baked) document, so both halves agree on heading levels.
            by_page = {}
            for group, detect in (
                ([i for i in indices if i not in scans], True),
                ([i for i in indices if i in scans], False),
            ):
                if not group:
                    continue
                part = pymupdf4llm.to_markdown(doc, pages = group, detect_bg_color = detect, **kwargs)
                if not isinstance(part, list) or len(part) != len(group):
                    return None
                by_page.update(zip(group, part))
            chunks = [by_page[i] for i in indices]
    except Exception:  # noqa: BLE001 - never let Markdown extraction break ingestion
        logger.warning("pymupdf4llm extraction failed; using plain text", exc_info = True)
        return None
    expected_pages = doc.page_count if pages is None else len(pages)
    if not isinstance(chunks, list) or len(chunks) != expected_pages:
        return None
    return [str(c.get("text") or "") for c in chunks]


def _pdf(
    source: str | bytes,
    want_images: bool,
    max_pages: int | None = None,
) -> tuple[list[Page], list[ParsedImage], int]:
    import fitz

    pages: list[Page] = []
    images: list[ParsedImage] = []
    doc = (
        fitz.open(stream = source, filetype = "pdf") if isinstance(source, bytes) else fitz.open(source)
    )
    try:
        if doc.needs_pass:
            raise ValueError("encrypted PDF requires a password")
        total_pages = doc.page_count
        page_numbers = range(total_pages if max_pages is None else min(total_pages, max_pages))
        if not config.PDF_MARKDOWN:
            md = None
        elif max_pages is None:
            md = _pdf_markdown(doc)
        else:
            md = _pdf_markdown(doc, page_numbers)
        for i, page_number in enumerate(page_numbers):
            page = doc[page_number]
            plain = page.get_text("text") or ""
            candidate = md[i] if md else ""
            # Prefer Markdown; use PyMuPDF text if Markdown is empty, mangles RTL/Indic, or loses most text.
            if (
                candidate
                and not _markdown_corrupted(candidate)
                and not _markdown_incomplete(candidate, plain)
            ):
                text = candidate
            else:
                text = plain
            # Scan detection uses the PDF, not Markdown: Markdown placeholders appear even without a text layer.
            images_on_page = page.get_image_info()
            needs_ocr = bool(images_on_page) and not plain.strip()
            for info in images_on_page:
                image_rect = fitz.Rect(info["bbox"]) & page.rect
                if image_rect.get_area() < page.rect.get_area() * 0.5:
                    continue
                # A selectable header/footer does not make the scanned body readable.
                body = fitz.Rect(
                    image_rect.x0,
                    image_rect.y0 + image_rect.height * 0.1,
                    image_rect.x1,
                    image_rect.y1 - image_rect.height * 0.1,
                )
                if len(page.get_text("text", clip = body).strip()) < config.OCR_MIN_CHARS:
                    needs_ocr = True
                    break
            if needs_ocr:
                text = plain
            pages.append(Page(text, page_number + 1, len(text), needs_ocr = needs_ocr))
            if want_images:
                for img in page.get_images(full = True):
                    xref = img[0]
                    try:
                        extracted = doc.extract_image(xref)
                    except Exception as exc:  # noqa: BLE001
                        logger.debug("skipping image xref %s: %s", xref, exc)
                        continue
                    image_bytes = extracted.get("image")
                    if image_bytes:
                        images.append(
                            ParsedImage(
                                image_bytes = image_bytes,
                                page_number = page_number + 1,
                                xref = xref,
                            )
                        )
    finally:
        doc.close()
    return pages, images, total_pages


def parse_pdf_bytes(data: bytes, *, max_pages: int | None = None) -> tuple[list[Page], int]:
    """Returns capped pages plus the full page count, so a max_pages cut differs from a short PDF."""
    pages, _images, total_pages = _pdf(data, want_images = False, max_pages = max_pages)
    return pages, total_pages


def mupdf_document_text(path: str) -> str:
    """Paragraph text of a reflowable e-book only MuPDF reads (.mobi, .fb2)."""
    import pymupdf

    with pymupdf.open(path) as doc:
        # A tall page keeps paragraphs from splitting at page breaks.
        doc.layout(width = 1000, height = 100000, fontsize = 11)
        paragraphs = [
            " ".join(block[4].split())
            for page in doc
            for block in page.get_text("blocks")
            if block[6] == 0
        ]
    return "\n\n".join(p for p in paragraphs if p)


def _merge_rects(boxes: list) -> list:
    """Union overlapping rectangles (largest-first) into figure regions."""
    import pymupdf

    rects = [pymupdf.Rect(b) for b in boxes]
    rects = [r for r in rects if r.width > 5 and r.height > 5]
    merged: list = []
    for box in sorted(rects, key = lambda r: -r.get_area()):
        placed = False
        for m in merged:
            if m.intersects(box):
                m |= box
                placed = True
                break
        if not placed:
            merged.append(+box)
    return merged


def _figure_boxes(
    page,
    *,
    min_area_frac: float = 0.04,
    min_side: float = 40.0,
) -> list:
    """Qualifying figure-region rectangles on a page: cluster vector drawings + raster
    placements, merge overlaps, keep the page-spanning ones (area/side filtered)."""
    boxes: list = []
    try:
        boxes.extend(info["bbox"] for info in page.get_image_info())
    except Exception:
        pass
    try:
        boxes.extend(page.cluster_drawings())
    except Exception:
        pass
    if not boxes:
        return []
    page_area = page.rect.width * page.rect.height
    keep: list = []
    for box in _merge_rects(boxes):
        if (
            box.get_area() >= min_area_frac * page_area
            and box.width >= min_side
            and box.height >= min_side
        ):
            keep.append(box)
    return keep


def pages_with_figures(
    path: str,
    *,
    max_pages: int = 4,
    min_area_frac: float = 0.04,
    min_side: float = 40.0,
    exclude_pages: set[int] | None = None,
) -> list[int]:
    """1-based pages with a figure region, for tiling; exclude_pages, already OCR'd whole, are skipped."""
    exclude = exclude_pages or set()
    try:
        import pymupdf
    except Exception:
        return []
    try:
        doc = pymupdf.open(path)
    except Exception:
        return []
    pages: list[int] = []
    try:
        for i, page in enumerate(doc):
            if (i + 1) in exclude:
                continue
            if _figure_boxes(page, min_area_frac = min_area_frac, min_side = min_side):
                pages.append(i + 1)
                if len(pages) >= max_pages:
                    break
        return pages
    finally:
        doc.close()


def render_pdf_figure_tiles(
    path: str,
    page_numbers,
    *,
    dpi: int = 200,
    rows: int = 2,
    cols: int = 2,
    overlap: float = 0.12,
    fullpage: bool = True,
    max_tiles: int = 24,
) -> list[ParsedImage]:
    """Render figure-bearing pages as overlapping high-DPI tiles (plus an optional full
    page), each a ``ParsedImage`` keyed by page number. Tiling keeps small labels legible
    and covers every sub-figure without exact region detection. Any failure yields []."""
    wanted = [int(n) for n in page_numbers]
    if not wanted:
        return []
    rows, cols = max(1, int(rows)), max(1, int(cols))
    try:
        import pymupdf
    except Exception:
        return []
    try:
        doc = pymupdf.open(path)
    except Exception:
        return []
    out: list[ParsedImage] = []
    try:
        for num in wanted:
            if num < 1 or num > doc.page_count:
                continue
            page = doc[num - 1]
            rect = page.rect
            clips: list = [rect] if fullpage else []
            cw, ch = rect.width / cols, rect.height / rows
            ox, oy = cw * overlap, ch * overlap
            for r in range(rows):
                for c in range(cols):
                    clips.append(
                        pymupdf.Rect(
                            rect.x0 + c * cw - ox,
                            rect.y0 + r * ch - oy,
                            rect.x0 + (c + 1) * cw + ox,
                            rect.y0 + (r + 1) * ch + oy,
                        )
                        & rect
                    )
            for index, clip in enumerate(clips):
                try:
                    pix = page.get_pixmap(dpi = dpi, clip = clip)
                    is_full_page = fullpage and index == 0
                    out.append(
                        ParsedImage(
                            image_bytes = pix.tobytes("png"),
                            page_number = num,
                            xref = 0,
                            full_page = is_full_page,
                            tile_index = None if is_full_page else index - int(fullpage),
                            tile_count = rows * cols,
                        )
                    )
                except Exception:
                    continue
                if len(out) >= max_tiles:
                    return out
        return out
    finally:
        doc.close()


def render_pdf_pages(
    path: str,
    page_numbers,
    *,
    dpi: int = 150,
) -> dict[int, bytes]:
    """Renders whole 1-based pages to PNG for scanned-page OCR; failures skip the page, never raise."""
    wanted = {int(n) for n in page_numbers}
    if not wanted:
        return {}
    try:
        import pymupdf
    except Exception:
        return {}
    try:
        doc = pymupdf.open(path)
    except Exception:
        return {}
    out: dict[int, bytes] = {}
    try:
        for i, page in enumerate(doc):
            num = i + 1
            if num not in wanted:
                continue
            try:
                pix = page.get_pixmap(dpi = dpi)
                out[num] = pix.tobytes("png")
            except Exception:
                continue
        return out
    finally:
        doc.close()


_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_MC_FALLBACK = "{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback"
_M = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"
_DOCX_MATH = frozenset((_M + "oMathPara", _M + "oMath"))
# Runs not shown as body text; text boxes are read as blocks of their own.
_DOCX_SKIP_RUNS_UNDER = frozenset(
    (_W + "del", _W + "moveFrom", _W + "rt", _W + "txbxContent", _MC_FALLBACK)
)
_DOCX_SKIP_RUNS_OR_MATH = _DOCX_SKIP_RUNS_UNDER | _DOCX_MATH
_DOCX_MATH_ROWS = {
    "oMathPara": ("oMath", "\n"),
    "eqArr": ("e", "\n"),
    "m": ("mr", " \\\\ "),
    "mr": ("e", " & "),
}
# Symbol-font characters from 0x20, as mammoth (the + attachment reader) maps them.
_DOCX_SYMBOL_FONTS = {
    "symbol": (
        " !∀#∃%&∍()*+,−./0123456789:;<=>?"
        "≅ΑΒΧΔΕΦΓΗΙϑΚΛΜΝΟΠΘΡΣΤΥςΩΞΨΖ[∴]⊥_"
        "‾αβχδεφγηιϕκλμνοπθρστυϖωξψζ{|}~\0"
        "\0\0\0\0\0\0\0\0\0\0\0\0\0\0\0\0\0\0\0\0\0\0\0\0\0\0\0\0\0\0\0\0"
        "€ϒ′≤⁄∞ƒ♣♦♥♠↔←↑→↓°±″≥×∝∂•÷≠≡≈…⏐⎯↵"
        "ℵℑℜ℘⊗⊕∅∩∪⊃⊇⊄⊂⊆∈∉∠∇®©™∏√⋅¬∧∨⇔⇐⇑⇒⇓"
        "◊〈®©™∑⎛⎜⎝⎡⎢⎣⎧⎨⎩⎪\0〉∫⌠⎮⌡⎞⎟⎠⎤⎥⎦⎫⎬⎭"
    ),
    "wingdings": (
        " 🖉✂✁👓🕭🕮🕯🕿✆🖂🖃📪📫📬📭🗀🗁🗎🗏🗐🗄⌛🖮🖰🖲🖳🖴🖫🖬✇✍"
        "🖎✌🖏👍👎☜☞☝☟🖐☺😐☹💣🕱🏳🏱✈☼🌢❄🕆✞🕈✠✡☪☯🕉☸♈♉"
        "♊♋♌♍♎♏♐♑♒♓🙰🙵⚫🔾◼🞏🞐❑❒🞟⧫◆❖⬩⌧⮹⌘🏵🏶🙶🙷▯"
        "🄋➀➁➂➃➄➅➆➇➈➉🄌➊➋➌➍➎➏➐➑➒➓🙢🙠🙡🙣🙞🙜🙝🙟∙•"
        "⬝⭘🞆🞈🞊🞋🔿▪🞎🟁🟅★🟋🟏🟓🟑⯐⌖⯎⯏⯑✪✰🕐🕑🕒🕓🕔🕕🕖🕗🕘"
        "🕙🕚🕛⮰⮱⮲⮳⮴⮵⮶⮷🙪🙫🙕🙔🙗🙖🙐🙑🙒🙓⌫⌦⮘⮚⮙⮛⮈⮊⮉⮋🡨"
        "🡪🡩🡫🡬🡭🡯🡮🡸🡺🡹🡻🡼🡽🡿🡾⇦⇨⇧⇩⬄⇳⬁⬀⬃⬂🢬🢭🗶✓🗷🗹"
    ),
    "wingdings 2": (
        " 🖊🖋🖌🖍✄✀🕾🕽🗅🗆🗇🗈🗉🗊🗋🗌🗍📋🗑🗔🖵🖶🖷🖸🖭🖯🖱🖒🖓🖘🖙"
        "🖚🖛👈👉🖜🖝🖞🖟🖠🖡👆👇🖢🖣🖑🗴🗸🗵☑⮽☒⮾⮿🛇⦸🙱🙴🙲🙳‽🙹🙺"
        "🙻🙦🙤🙥🙧🙚🙘🙙🙛⓪①②③④⑤⑥⑦⑧⑨⑩⓿❶❷❸❹❺❻❼❽❾❿\0"
        "☉🌕☽☾⸿✝🕇🕜🕝🕞🕟🕠🕡🕢🕣🕤🕥🕦🕧🙨🙩⋅🞄⦁●○🞅🞇🞉⊙⦿🞌"
        "🞍◾■□🞑🞒🞓🞔▣🞕🞖🞗🞘⬩⬥◇🞚◈🞛🞜🞝🞞⬪⬧◊🞠◖◗⯊⯋⯀⯁"
        "⬟⯂⬣⬢⯃⯄🞡🞢🞣🞤🞥🞦🞧🞨🞩🞪🞫🞬🞭🞮🞯🞰🞱🞲🞳🞴🞵🞶🞷🞸🞹🞺"
        "🞻🞼🞽🞾🞿🟀🟂🟄🟆🟉🟊✶🟌🟎🟐🟒✹🟃🟇✯🟍🟔⯌⯍※⁂"
    ),
    "wingdings 3": (
        " ⭠⭢⭡⭣⭦⭧⭩⭨⭰⭲⭱⭳⭶⭸⭻⭽⭤⭥⭪⭬⭫⭭⭍⮠⮡⮢⮣⮤⮥⮦⮧"
        "⮐⮑⮒⮓⮀⮃⭾⭿⮄⮆⮅⮇⮏⮍⮎⮌⭮⭯⎋⌤⌃⌥␣⍽⇪⮸🢠🢡🢢🢣🢤🢥"
        "🢦🢧🢨🢩🢪🢫🡐🡒🡑🡓🡔🡕🡗🡖🡘🡙▲▼△▽◀▶◁▷◣◢◤◥🞀🞂🞁\0"
        "🞃⯅⯆⯇⯈⮜⮞⮝⮟🠐🠒🠑🠓🠔🠖🠕🠗🠘🠚🠙🠛🠜🠞🠝🠟🠀🠂🠁🠃🠄🠆🠅"
        "🠇🠈🠊🠉🠋🠠🠢🠤🠦🠨🠪🠬🢜🢝🢞🢟🠮🠰🠲🠴🠶🠸🠺🠹🠻🢘🢚🢙🢛🠼🠾🠽"
        "🠿🡀🡂🡁🡃🡄🡆🡅🡇⮨⮩⮪⮫⮬⮭⮮⮯🡠🡢🡡🡣🡤🡥🡧🡦🡰🡲🡱🡳🡴🡵🡷"
        "🡶🢀🢂🢁🢃🢄🢅🢇🢆🢐🢒🢑🢓🢔🢖🢕🢗"
    ),
    "webdings": (
        " 🕷🕸🕲🕶🏆🎖🖇🗨🗩🗰🗱🌶🎗🙾🙼🗕🗖🗗⏴⏵⏶⏷⏪⏩⏮⏭⏸⏹⏺🗚🗳"
        "🛠🏗🏘🏙🏚🏜🏭🏛🏠🏖🏝🛣🔍🏔👁👂🏞🏕🛤🏟🛳🕬🕫🕨🔈🎔🎕🗬🙽🗭🗪🗫"
        "⮔✔🚲⬜🛡📦🛱⬛🚑🛈🛩🛰🟈🕴⬤🛥🚔🗘🗙❓🛲🚇🚍⛳⦸⊖🚭🗮⏐🗯🗲\0"
        "🚹🚺🛉🛊🚼👽🏋⛷🏂🏌🏊🏄🏍🏎🚘🗠🛢💰🏷💳👪🗡🗢🗣✯🖄🖅🖃🖆🖹🖺🖻"
        "🕵🕰🖽🖾📋🗒🗓🕮📚🗞🗟🗃🗂🖼🎭🎜🎘🎙🎧💿🎞📷🎟🎬📽📹📾📻🎚🎛📺💻"
        "🖥🖦🖧🕹🎮🕻🕼📟🖁🖀🖨🖩🖿🖪🗜🔒🔓🗝📥📤🕳🌣🌤🌥🌦☁🌨🌧🌩🌪🌬🌫"
        "🌜🌡🛋🛏🍽🍸🛎🛍Ⓟ♿🛆🖈🎓🗤🗥🗦🗧🛪🐿🐦🐟🐕🐈🙬🙮🙭🙯🗺🌍🌏🌎🕊"
    ),
}


def _docx_placeholder(element) -> bool:
    # An unfilled content control holds Word's prompt ("Click or tap here to enter text."), not a value.
    if element.tag != _W + "sdt":
        return False
    flag = element.find(_W + "sdtPr/" + _W + "showingPlcHdr")
    return flag is not None and flag.get(_W + "val", "true") not in ("0", "false", "off")


def _docx_inside(element, stop, tags) -> bool:
    node = element.getparent()
    while node is not stop:
        if node.tag in tags or _docx_placeholder(node):
            return True
        node = node.getparent()
    return False


def _docx_unwrap_table_controls(body) -> None:
    # python-docx skips w:tr / w:tc wrapped in content controls (cover pages, repeating sections).
    for wrapper in list(body.iter(_W + "sdt", _W + "customXml")):
        parent = wrapper.getparent()
        if parent is None or parent.tag not in (_W + "tbl", _W + "tr"):
            continue
        content = wrapper.find(_W + "sdtContent") if wrapper.tag == _W + "sdt" else wrapper
        keep = (_W + "tr", _W + "tc", _W + "sdt", _W + "customXml")
        placeholder = _docx_placeholder(wrapper)
        idx = parent.index(wrapper)
        for i, child in enumerate(
            [c for c in (content if content is not None else ()) if c.tag in keep]
        ):
            if placeholder:  # keep the cells so columns line up, drop the prompt text
                for tc in child.iter(_W + "tc"):
                    for el in [e for e in tc if e.tag != _W + "tcPr"]:
                        tc.remove(el)
            parent.insert(idx + i, child)
        parent.remove(wrapper)


def _docx_blocks(element, parent):
    """Paragraphs and tables in document order, including content controls and text boxes."""
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    for child in element:
        if child.tag == _W + "p":
            yield Paragraph(child, parent)
            for box in child.iter(_W + "txbxContent"):
                # Nested boxes are read with their outer box; fallback copies repeat the real one.
                if not _docx_inside(box, child, _DOCX_SKIP_RUNS_UNDER):
                    yield from _docx_blocks(box, parent)
        elif child.tag == _W + "tbl":
            yield Table(child, parent)
        elif child.tag in (_W + "sdt", _W + "customXml") and not _docx_placeholder(child):
            content = child.find(_W + "sdtContent")
            yield from _docx_blocks(child if content is None else content, parent)


def _docx_math_text(element) -> str:
    # Mirrored by docxMathText in the frontend's attachment-content.ts.
    tag = element.tag
    if tag in _DOCX_SKIP_RUNS_UNDER or _docx_placeholder(element):
        return ""
    if tag == _W + "r":
        return _docx_run_text(element)
    if tag == _W + "sym":
        return _docx_symbol(element)
    if tag == _M + "t":
        return element.text or ""
    name = tag[len(_M) :] if tag.startswith(_M) else ""

    def arg(key):
        child = element.find(_M + key)
        if child is None or prop(f"{key}Hide", "off") not in ("0", "false", "off"):
            return ""
        return _docx_math_text(child)

    def prop(key, default):
        node = element.find(f"{_M}{name}Pr/{_M}{key}")
        return default if node is None else node.get(_M + "val", "")

    def scripts(sub, sup):
        return (f"_{{{sub}}}" if sub else "") + (f"^{{{sup}}}" if sup else "")

    if name == "f":
        if prop("type", "bar") == "noBar":
            return f"{{{arg('num')} \\atop {arg('den')}}}"
        return f"\\frac{{{arg('num')}}}{{{arg('den')}}}"
    if name == "phant" and prop("show", "on") in ("0", "false", "off"):
        return ""
    if name in ("sSub", "sSup", "sSubSup"):
        return arg("e") + scripts(arg("sub"), arg("sup"))
    if name == "sPre":
        return "{}" + scripts(arg("sub"), arg("sup")) + arg("e")
    if name == "limLow":
        return arg("e") + scripts(arg("lim"), "")
    if name == "limUpp":
        return arg("e") + scripts("", arg("lim"))
    if name == "nary":
        return prop("chr", "\u222b") + scripts(arg("sub"), arg("sup")) + arg("e")
    if name == "rad":
        deg = arg("deg")
        return f"\\sqrt[{deg}]{{{arg('e')}}}" if deg else f"\\sqrt{{{arg('e')}}}"
    if name == "acc":
        return arg("e") + prop("chr", "\u0302")
    if name in ("bar", "groupChr"):
        side = "over" if prop("pos", "bot") == "top" else "under"
        if name == "bar":
            return f"\\{side}line{{{arg('e')}}}"
        mark = prop("chr", "\u23df")
        if mark in ("\u23de", "\u23df"):
            return f"\\{side}brace{{{arg('e')}}}"
        return f"\\{side}set{{{mark}}}{{{arg('e')}}}"
    if name == "func":
        return f"{arg('fName')} {arg('e')}"
    if name == "d":
        return (
            prop("begChr", "(")
            + prop("sepChr", "|").join(_docx_math_text(e) for e in element.iterchildren(_M + "e"))
            + prop("endChr", ")")
        )
    if name in _DOCX_MATH_ROWS:
        child, sep = _DOCX_MATH_ROWS[name]
        return sep.join(_docx_math_text(c) for c in element.iterchildren(_M + child))
    return "".join(_docx_math_text(child) for child in element.iterchildren("*"))


def _docx_checkbox(field) -> str:
    box = field.find(f"{_W}ffData/{_W}checkBox")
    if box is None:
        return ""
    flag = box.find(_W + "checked")
    if flag is None:
        flag = box.find(_W + "default")
    ticked = flag is not None and flag.get(_W + "val", "true") not in ("0", "false", "off")
    return "☒" if ticked else "☐"


def _docx_symbol(sym) -> str:
    try:
        code = int(sym.get(_W + "char") or "", 16)
    except ValueError:
        return ""
    table = _DOCX_SYMBOL_FONTS.get((sym.get(_W + "font") or "").lower())
    if table is not None and code >> 8 in (0, 0xF0):
        index = (code & 0xFF) - 0x20
        code = ord(table[index]) if 0 <= index < len(table) else 0
    return chr(code) if 0 <= code <= 0x10FFFF and chr(code).isprintable() else ""


def _docx_run_text(run) -> str:
    if run.find(f"{_W}fldChar/{_W}ffData/{_W}checkBox") is None and run.find(_W + "sym") is None:
        return run.text
    return "".join(
        _docx_checkbox(e)
        if e.tag == _W + "fldChar"
        else _docx_symbol(e)
        if e.tag == _W + "sym"
        else str(e)
        for e in run.xpath(
            "w:br | w:cr | w:noBreakHyphen | w:ptab | w:t | w:tab | w:fldChar | w:sym"
        )
    )


def _docx_paragraph_text(paragraph) -> str:
    """Paragraph.text skips runs wrapped in w:ins, w:sdt, w:fldSimple, w:smartTag, and equations."""
    p = paragraph._p
    return "".join(
        _docx_math_text(node) if node.tag in _DOCX_MATH else _docx_run_text(node)
        for node in p.iter(_W + "r", *_DOCX_MATH)
        if not _docx_inside(node, p, _DOCX_SKIP_RUNS_OR_MATH)
    )


def _docx_int(value) -> int:
    try:
        return max(-(10**18), min(int(value or 0), 10**18))
    except ValueError:
        return 0


def _docx_number_lists(document) -> None:
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement, parse_xml

    def related(reltype):
        part = next(
            (
                r.target_part
                for r in document.part.rels.values()
                if r.reltype == reltype and not r.is_external
            ),
            None,
        )
        if part is None:
            return None
        return part.element if hasattr(part, "element") else parse_xml(part.blob)

    def val(node, name):
        child = None if node is None else node.find(_W + name)
        return None if child is None else child.get(_W + "val") or None

    def first_by_level(nodes):
        by_level = {}
        for node in nodes:
            by_level.setdefault(node.get(_W + "ilvl"), node)
        return by_level

    numbering = related(RT.NUMBERING)
    if numbering is None:
        return
    styles = related(RT.STYLES)
    style_by_id = {
        style.get(_W + "styleId"): style
        for style in (() if styles is None else styles.iterchildren(_W + "style"))
    }
    default_style = next(
        (
            style_id
            for style_id, style in style_by_id.items()
            if style.get(_W + "type") == "paragraph"
            and style.get(_W + "default") in ("1", "true", "on")
        ),
        None,
    )
    abstracts = {a.get(_W + "abstractNumId"): a for a in numbering.iterchildren(_W + "abstractNum")}
    nums = {n.get(_W + "numId"): n for n in numbering.iterchildren(_W + "num")}

    # numId and ilvl inherit through basedOn independently, like any other paragraph property.
    def style_num(
        style_id,
        name,
        depth = 0,
    ):
        style = style_by_id.get(style_id)
        if style is None or depth > 20:
            return None
        value = val(style.find(f"{_W}pPr/{_W}numPr"), name)
        return style_num(val(style, "basedOn"), name, depth + 1) if value is None else value

    def linked(abstract_id, depth = 0):
        num_id = style_num(val(abstracts.get(abstract_id), "numStyleLink"), "numId")
        target = val(nums.get(num_id), "abstractNumId")
        if target in (None, abstract_id) or depth > 5:
            return abstract_id
        return linked(target, depth + 1)

    instances = {}

    def instance(num_id):
        if num_id in instances:
            return instances[num_id]
        num = nums.get(num_id)
        abstract_id = linked(val(num, "abstractNumId"))
        abstract = abstracts.get(abstract_id)
        instances[num_id] = None
        if num is None or abstract is None:
            return None
        overrides = first_by_level(num.iterchildren(_W + "lvlOverride"))
        defined = first_by_level(abstract.iterchildren(_W + "lvl"))
        levels = []
        for index in range(9):
            override = overrides.get(str(index))
            lvl = None if override is None else override.find(_W + "lvl")
            lvl = defined.get(str(index)) if lvl is None else lvl
            legal = lvl is not None and any(
                node.get(_W + "val", "true") not in ("0", "false", "off")
                for node in lvl.iterchildren(_W + "isLgl")
            )
            start = _docx_int(val(override, "startOverride") or val(lvl, "start"))
            fmt, restart, text = (val(lvl, key) for key in ("numFmt", "lvlRestart", "lvlText"))
            levels.append((lvl, start, fmt, restart, legal, text or ""))
        restarts = [
            int(at)
            for at, node in overrides.items()
            if at in tuple("012345678") and node.find(_W + "startOverride") is not None
        ]
        by_style = {}
        for index, (lvl, *_) in enumerate(levels):
            by_style.setdefault(val(lvl, "pStyle"), index)
        instances[num_id] = (abstract_id, levels, restarts, by_style)
        return instances[num_id]

    w15 = "{http://schemas.microsoft.com/office/word/2012/wordml}"
    restart_after_break = {
        linked(abstract_id)
        for abstract_id, abstract in abstracts.items()
        if abstract.get(w15 + "restartNumberingAfterBreak") in ("1", "true", "on")
    }
    counters: dict[str, list] = {}
    started: set[str] = set()
    restartable: set[str] = set()

    def label(p_pr):
        mark = None if p_pr is None else p_pr.find(_W + "rPr")
        if mark is not None and (
            mark.find(_W + "del") is not None or mark.find(_W + "moveFrom") is not None
        ):
            return None
        direct = None if p_pr is None else p_pr.find(_W + "numPr")
        style = style_id = val(p_pr, "pStyle") or default_style
        num_id = val(direct, "numId") or style_num(style, "numId")
        found = None if num_id is None else instance(num_id)
        if found is None:
            return None
        abstract_id, levels, restarts, by_style = found
        # A heading style's numPr often names only the list, not the level.
        linked_level = None
        for _ in range(21):
            if style is None or linked_level is not None:
                break
            linked_level = by_style.get(style)
            style = val(style_by_id.get(style), "basedOn")
        ilvl = val(direct, "ilvl") or style_num(style_id, "ilvl") or linked_level
        ilvl = min(8, max(0, _docx_int(ilvl)))
        lvl, _, fmt, _, legal, text = levels[ilvl]
        if lvl is None:
            return None
        # Word shares counters across instances of one abstract; a startOverride restarts them once.
        counts = counters.setdefault(abstract_id, [None] * 9)
        if abstract_id in restart_after_break:
            restartable.add(abstract_id)
        if num_id not in started:
            started.add(num_id)
            for at in restarts:
                counts[at:] = [None] * (9 - at)
        for i in range(ilvl):
            if counts[i] is None:
                counts[i] = levels[i][1]
        counts[ilvl] = levels[ilvl][1] if counts[ilvl] is None else counts[ilvl] + 1
        for i in range(ilvl + 1, 9):
            restart = levels[i][3]
            if restart is None or ilvl < _docx_int(restart):
                counts[i] = None
        if fmt == "bullet" or len(text) > 256:
            return None
        value = ""
        for index, piece in enumerate(re.split("%([1-9])", text)):
            if index % 2 == 0:
                value += piece
            else:
                _, start, piece_fmt, *_ = levels[int(piece) - 1]
                if legal and piece_fmt != "none" and not (piece_fmt or "").startswith("decimal"):
                    piece_fmt = "decimal"
                count = counts[int(piece) - 1]
                value += _list_number(start if count is None else count, piece_fmt)
            if len(value) > 256:
                return None
        return value if value.strip() else None

    body = document.element.body
    hidden = _DOCX_SKIP_RUNS_UNDER - {_W + "txbxContent"}
    labels = []
    budget = 1 << 20
    for p in body.iter(_W + "p"):
        if _docx_inside(p, body, hidden):
            # Word still counts a numbered paragraph that only holds placeholder text.
            if not any(node.tag in hidden for node in p.iterancestors()):
                label(p.find(_W + "pPr"))
            continue
        p_pr = p.find(_W + "pPr")
        value = label(p_pr)
        if value is not None:
            labels.append((p, p_pr, value))
            budget -= len(value)
            if budget < 0:
                return
        if p_pr is not None and restartable and p_pr.find(_W + "sectPr") is not None:
            for abstract_id in restartable:
                counters.pop(abstract_id, None)
            restartable.clear()
    for p, p_pr, value in labels:
        text = OxmlElement("w:t")
        text.text = f"{value} "
        run = OxmlElement("w:r")
        run.append(text)
        if p_pr is None:
            p.insert(0, run)
        else:
            p_pr.addnext(run)


def _docx_table_rows(table) -> list[str]:
    """Cells joined by pipes, aligned to the layout grid; the locator splits anchors on pipes."""
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    rows: list[str] = []
    seen: set = set()
    for row in table.rows:
        cells: list[str] = [""] * getattr(row, "grid_cols_before", 0)
        trailing: list[str] = []
        for cell in row.cells:
            # Merged cells share one <w:tc>: emit text once, then placeholders to keep alignment.
            if cell._tc in seen:
                cells.append("")
                continue
            seen.add(cell._tc)
            # Text before a nested table is the aligned field; the table and anything after it flatten below.
            field: list[str] = []
            after_table = False
            for item in _docx_blocks(cell._tc, cell):
                if isinstance(item, Table):
                    after_table = True
                    trailing.extend(_docx_table_rows(item))
                elif isinstance(item, Paragraph):
                    text = " ".join(_docx_paragraph_text(item).split())
                    if text:
                        (trailing if after_table else field).append(text)
            cells.append(" ".join(field))  # empty cells kept so columns line up
        cells.extend([""] * getattr(row, "grid_cols_after", 0))
        if any(c.strip() for c in cells):
            rows.append(" | ".join(cells))
        rows.extend(trailing)
    return rows


_WORD_EXTS = frozenset((".docx", ".docm", ".dotx", ".dotm"))
_WORD_MAIN_TYPES = (
    "application/vnd.ms-word.document.macroEnabled.main+xml",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml",
    "application/vnd.ms-word.template.macroEnabledTemplate.main+xml",
)


def _open_word(path: str):
    # docx.Document() rejects the macro and template main parts.
    import docx  # noqa: F401  (registers the part types)
    from docx.opc.constants import CONTENT_TYPE as CT
    from docx.opc.part import PartFactory
    from docx.package import Package
    from docx.parts.document import DocumentPart

    for content_type in _WORD_MAIN_TYPES:
        PartFactory.part_type_for.setdefault(content_type, DocumentPart)
    part = Package.open(path).main_document_part
    if part.content_type not in (CT.WML_DOCUMENT_MAIN, *_WORD_MAIN_TYPES):
        raise ValueError(
            f"file '{os.path.basename(path)}' is not a Word file, content type is '{part.content_type}'"
        )
    return part.document


def _docx(path: str) -> list[Page]:
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = _open_word(path)
    lines: list[str] = []
    _docx_unwrap_table_controls(document.element.body)
    try:
        _docx_number_lists(document)
    except Exception as exc:
        logger.warning("Could not number lists in %s: %s", os.path.basename(path), exc)
    label_notes = _docx_mark_notes(document)
    # Walk body content in document order: paragraphs alone drop tables entirely.
    for block in _docx_blocks(document.element.body, document):
        if isinstance(block, Paragraph):
            text = _docx_paragraph_text(block)
            if text.strip():
                lines.append(text)
        elif isinstance(block, Table):
            lines.extend(_docx_table_rows(block))
    return [_page(label_notes("\n".join(lines)), None)]


def _roman(n: int) -> str:
    out = ""
    for value, digits in zip(
        (1000, 900, 500, 400, 100, 90, 50, 40, 10, 9, 5, 4, 1),
        ("m", "cm", "d", "cd", "c", "xc", "l", "xl", "x", "ix", "v", "iv", "i"),
    ):
        count, n = divmod(n, value)
        out += digits * count
    return out


def _list_number(n: int, fmt: str | None) -> str:
    fmt = fmt or ""
    if fmt == "none":
        return ""
    if fmt == "ordinal" and n >= 1:
        suffix = "th" if 11 <= n % 100 <= 13 else {1: "st", 2: "nd", 3: "rd"}.get(n % 10, "th")
        return f"{n}{suffix}"
    lower = fmt.startswith("lower")
    limit = 3999 if fmt.endswith("Roman") else 26 * 256 if fmt.endswith("Letter") else n
    if not (lower or fmt.startswith("upper")) or not 1 <= n <= limit:
        return f"0{n}" if fmt == "decimalZero" and 0 <= n < 10 else str(n)
    if fmt.endswith("Roman"):
        text = _roman(n)
    elif fmt.endswith("Alpha"):
        # CSS lower-alpha is bijective (z, aa, ab); Word's lowerLetter repeats (z, aa, bb).
        text = ""
        while n > 0:
            n, rest = divmod(n - 1, 26)
            text = chr(97 + rest) + text
    else:
        text = chr(97 + (n - 1) % 26) * ((n + 25) // 26)
    return text if lower else text.upper()


def _docx_mark_notes(document):
    """Sentinels after body note references; the returned function labels the surviving ones (1, 2 / i, ii) and appends the notes."""
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement, parse_xml
    from docx.table import Table

    kinds = []
    for kind, reltype, label in (
        ("footnote", RT.FOOTNOTES, str),
        ("endnote", RT.ENDNOTES, _roman),
    ):
        part = next(
            (
                r.target_part
                for r in document.part.rels.values()
                if r.reltype == reltype and not r.is_external
            ),
            None,
        )
        if part is None:
            continue
        bodies: dict[str, str] = {}
        for note in parse_xml(part.blob).iterchildren(_W + kind):
            if note.get(_W + "type", "normal") != "normal":
                continue
            texts = []
            for block in _docx_blocks(note, document):
                if isinstance(block, Table):
                    texts.extend(_docx_table_rows(block))
                else:
                    texts.append(" ".join(_docx_paragraph_text(block).split()))
            bodies[note.get(_W + "id")] = " ".join(t for t in texts if t)
        if bodies:
            kinds.append((kind, label, bodies))
    if not kinds:
        return lambda text: text

    refs: list[tuple[int, str]] = []
    # Nonce: document text shaped like a sentinel stays as written.
    nonce = secrets.token_hex(4)
    sentinel = re.compile(f"\ue000{nonce}\\.(\\d+)\ue001")
    referenced: list[set[str]] = [set() for _ in kinds]
    for k, (kind, _, bodies) in enumerate(kinds):
        for ref in document.element.body.iter(_W + kind + "Reference"):
            note_id = ref.get(_W + "id")
            if note_id in bodies:
                referenced[k].add(note_id)
                marker = OxmlElement("w:t")
                marker.text = f"\ue000{nonce}.{len(refs)}\ue001"
                refs.append((k, note_id))
                ref.addnext(marker)

    def label_notes(text: str) -> str:
        numbers: list[dict[str, int]] = [{} for _ in kinds]

        def label(match) -> str:
            index = int(match.group(1))
            if index >= len(refs):
                return match.group(0)
            k, note_id = refs[index]
            number = numbers[k].setdefault(note_id, len(numbers[k]) + 1)
            return f"[{kinds[k][1](number)}]"

        text = sentinel.sub(label, text)
        lines = [text] if text else []
        for k, (kind, label_of, bodies) in enumerate(kinds):
            # Unreferenced notes stay; ones referenced only from deleted or moved text go.
            for note_id in bodies:
                if note_id not in referenced[k]:
                    numbers[k].setdefault(note_id, len(numbers[k]) + 1)
            notes = [
                f"[{label_of(number)}] {bodies[note_id]}"
                for note_id, number in sorted(numbers[k].items(), key = lambda item: item[1])
                if bodies[note_id]
            ]
            if notes:
                lines += [kind.capitalize() + "s", *notes]
        return "\n".join(lines)

    return label_notes


_HIGH_BYTES = bytes(range(0x80, 0x100))
_ASCII_LETTERS = b"ABCDEFGHIJKLMNOPQRSTUVWXYZabcdefghijklmnopqrstuvwxyz"
_HIGH_RUN = re.compile(rb"[\x80-\xff]+")


def _declared_charset(data: bytes) -> str | None:
    # Lazy: tools is heavy, and only HTML that is not UTF-8 gets here.
    from ..inference.tools import _META_CHARSET_SCAN_BYTES, _sniff_meta_charset
    return _sniff_meta_charset(data[:_META_CHARSET_SCAN_BYTES], "text/html")


def _decode_text(data: bytes, *, html: bool = False) -> str:
    # Check UTF-32 before its overlapping UTF-16 prefix.
    for bom, codec in (
        (codecs.BOM_UTF32_LE, "utf-32"),
        (codecs.BOM_UTF32_BE, "utf-32"),
        (codecs.BOM_UTF16_LE, "utf-16"),
        (codecs.BOM_UTF16_BE, "utf-16"),
        (codecs.BOM_UTF8, "utf-8-sig"),
    ):
        if data.startswith(bom):
            return data.decode(codec, errors = "replace")
    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError:
        pass
    else:
        # ISO-2022-JP is 7-bit, so it always passes as UTF-8; its escapes give it away.
        if html and "\x1b$" in text and _declared_charset(data) == "iso2022_jp":
            return data.decode("iso2022_jp", errors = "replace")
        return text
    declared = _declared_charset(data) if html else None
    # WHATWG reads UTF-16 labels as UTF-8, which these bytes already failed.
    if declared and declared != "utf-8":
        return data.decode(declared, errors = "replace")
    text = data.decode("utf-8", errors = "replace")
    # Ties stay UTF-8 (truncated file); cp1252 can form a stray valid sequence ("à\xa0»").
    non_ascii = len(text) - len(text.encode("ascii", "ignore"))
    if non_ascii >= 2 * text.count("\ufffd"):
        return text
    high = len(data) - len(data.translate(None, _HIGH_BYTES))
    letters = len(data) - len(data.translate(None, _ASCII_LETTERS))
    # A Latin-alphabet text never has half as many accented letters as plain ones; a few bytes say nothing.
    if high >= 8 and 2 * high > letters:
        from charset_normalizer import from_bytes

        legacy = [
            "cp1252",
            "gb18030",
            "cp950",
            "cp932",
            "cp949",
            "cp1251",
            "cp1253",
            "cp1255",
            "cp1256",
        ]
        results = from_bytes(data, cp_isolation = legacy)
        match = results.best()
        if match is not None:
            guess = str(match)
            # Ties are ambiguous; need high-byte words, not lone accents; unpaired CJK bytes mean katakana.
            tied = any(
                other is not match
                and (other.chaos, other.coherence) == (match.chaos, match.coherence)
                for other in results
            )
            if match.encoding == "cp1252":
                plausible = True
            elif match.encoding in ("cp1251", "cp1253", "cp1255", "cp1256"):
                in_words = sum(len(run) for run in _HIGH_RUN.findall(data) if len(run) >= 3)
                plausible = match.coherence > 0 and 2 * in_words > high
            else:
                plausible = len(guess) < len(data)
            if not tied and plausible:
                return guess
    return data.decode("cp1252", errors = "replace")


def parse(path: str, *, want_images: bool = False):
    """Parse a file into pages by extension. Returns ``list[Page]``, or
    ``(list[Page], list[ParsedImage])`` when ``want_images=True`` (only PDFs yield
    images). Raises ValueError on unsupported ext."""
    ext = os.path.splitext(path)[1].lower()

    if ext == ".pdf":
        pages, images, _total = _pdf(path, want_images)
        return (pages, images) if want_images else pages

    if ext in _WORD_EXTS:
        pages = _docx(path)
        return (pages, []) if want_images else pages

    if ext in config.DOCUMENT_UPLOAD_EXTS:
        pages = _office_document(path, ext)
        return (pages, []) if want_images else pages

    # .dot/.pot are Graphviz and gettext text too; only an OLE compound file is a Word/PowerPoint template.
    if ext in _OFFICE_TEMPLATE_EXTS and _is_compound_file(path):
        pages = _office_document(path, _OFFICE_TEMPLATE_EXTS[ext])
        return (pages, []) if want_images else pages

    if ext in (".html", ".htm", ".txt", ".md", ".markdown") or ext in config.SOURCE_TEXT_EXTS:
        is_html = ext in (".html", ".htm")
        with open(path, "rb") as f:
            raw = _decode_text(f.read(), html = is_html)
        # NUL never appears in source text; catches binaries under source extensions.
        if ext in config.SOURCE_TEXT_EXTS and "\x00" in raw:
            raise ValueError(f"unsupported binary content in text file: {os.path.basename(path)}")
        # Universal newlines, as text-mode open() gave: the chunker splits on "\n\n".
        raw = raw.replace("\r\n", "\n").replace("\r", "\n")
        pages = _html(raw) if is_html else [_page(raw, None)]
        return (pages, []) if want_images else pages

    raise ValueError(f"unsupported file type: {ext}")


_OFFICE_TEMPLATE_EXTS = {".dot": ".doc", ".pot": ".ppt"}
_OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


def _is_compound_file(path: str) -> bool:
    with open(path, "rb") as f:
        return f.read(len(_OLE_MAGIC)) == _OLE_MAGIC


def _html_bytes_text(raw: bytes) -> str:
    return "\n\n".join(page.text for page in _html(_decode_text(raw, html = True)))


def _xhtml(path: str) -> list[tuple[str, int | None]]:
    with open(path, "rb") as f:
        return [(_html_bytes_text(f.read()), None)]


def _office_document(path: str, ext: str) -> list[Page]:
    from . import office_formats as office

    readers = {
        ".doc": office.doc,
        **dict.fromkeys((".xls", ".xlt"), office.xls),
        **dict.fromkeys((".ppt", ".pps"), office.ppt),
        ".rtf": office.rtf,
        ".epub": lambda p: office.epub(p, _html_bytes_text),
        ".msg": lambda p: office.msg(p, _html_bytes_text),
        **dict.fromkeys((".xlsx", ".xlsm", ".xltx", ".xltm"), office.xlsx),
        **dict.fromkeys((".pptx", ".pptm", ".potx", ".potm", ".ppsx", ".ppsm"), office.pptx),
        **dict.fromkeys((".odt", ".ods", ".odp", ".ott", ".ots", ".otp"), office.opendocument),
        # .mht/.mhtml are MIME with an HTML body.
        **dict.fromkeys((".eml", ".mht", ".mhtml"), lambda p: office.eml(p, _html_bytes_text)),
        **dict.fromkeys((".xhtml", ".xht"), _xhtml),
    }
    try:
        sections = readers[ext](path)
    except (KeyError, IndexError, struct.error, UnicodeError, EOFError) as exc:
        raise ValueError(f"could not read {os.path.basename(path)}: file is damaged") from exc
    return [_page(text, number) for text, number in sections if text.strip()]


def parse_text(text: str) -> list[Page]:
    """Wrap already-extracted text as a single Page (tests / in-memory ingest)."""
    return [_page(text, None)]
